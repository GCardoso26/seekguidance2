#!/usr/bin/env python3
"""Gera Visio (.vsdx) + Excel (Data Visualizer) das consultas do onboarding.

Fonte: docs/engineering/TEAM_ONBOARDING_ESTADO_ATUAL.md
Stencil: Judge_doc.vsdx extraído em /tmp/visio_vsdx
"""
from __future__ import annotations

import html
import shutil
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "architecture" / "visio"
TEMPLATE = Path("/tmp/visio_vsdx")
NS = "http://schemas.microsoft.com/office/visio/2012/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Masters do Judge_doc.vsdx
M_START, M_DATA, M_DECISION, M_PROCESS = 2, 4, 5, 6
M_CONN, M_ROUND, M_GRAD, M_SUB, M_EXT = 7, 9, 13, 16, 17

MASTER_RELS = [
    (1, "master1.xml"),
    (2, "master2.xml"),
    (3, "master3.xml"),
    (4, "master4.xml"),
    (5, "master5.xml"),
    (6, "master6.xml"),
    (7, "master7.xml"),
    (8, "master8.xml"),
    (9, "master9.xml"),
    (10, "master10.xml"),
]


@dataclass
class Node:
    key: str
    text: str
    kind: str  # start, process, sub, decision, data, ext, end, note
    x: float
    y: float
    w: float = 2.05
    h: float = 0.62
    href: str | None = None


@dataclass
class Edge:
    src: str
    dst: str
    label: str = ""
    color: str = "#1998d6"  # azul; #75d175 sim; #e51400 não


@dataclass
class Page:
    name: str
    name_u: str
    width: float
    height: float
    nodes: list[Node]
    edges: list[Edge]
    notes: list[str] = field(default_factory=list)


def uid() -> str:
    return "{" + str(uuid.uuid4()).upper() + "}"


def master_for(kind: str) -> int:
    return {
        "start": M_START,
        "end": M_START,
        "process": M_PROCESS,
        "sub": M_SUB,
        "decision": M_DECISION,
        "data": M_DATA,
        "ext": M_EXT,
        "note": M_ROUND,
    }[kind]


def name_for(kind: str) -> tuple[str, str]:
    return {
        "start": ("Start/End", "Início/Término"),
        "end": ("Start/End", "Início/Término"),
        "process": ("Process", "Processo"),
        "sub": ("Subprocess", "Subprocesso"),
        "decision": ("Decision", "Decisão"),
        "data": ("Data", "Dados"),
        "ext": ("External Data", "Dados Externos"),
        "note": ("Rounded Rectangle", "Retângulo Arredondado"),
    }[kind]


def shape_xml(sid: int, node: Node) -> str:
    mid = master_for(node.kind)
    nu, nn = name_for(node.kind)
    layer = "0"
    txt = escape(node.text)
    extra = ""
    if node.href:
        extra = (
            f"<Section N='Hyperlink'><Row N='Row_1'>"
            f"<Cell N='Description' V=''/><Cell N='Address' V=''/>"
            f"<Cell N='SubAddress' V='{escape(node.href)}' F='Pages[{escape(node.href)}]!ThePage!PAGENAME()'/>"
            f"<Cell N='ExtraInfo' V=''/><Cell N='Frame' V=''/>"
            f"<Cell N='NewWindow' V='0'/><Cell N='Default' V='0'/>"
            f"<Cell N='Invisible' V='0'/><Cell N='SortKey' V=''/>"
            f"</Row></Section>"
        )
    w = node.w
    h = node.h
    return (
        f"<Shape ID='{sid}' NameU='{nu}.{sid}' Name='{nn}.{sid}' Type='Shape' Master='{mid}'>"
        f"<Cell N='PinX' V='{node.x}'/>"
        f"<Cell N='PinY' V='{node.y}'/>"
        f"<Cell N='Width' V='{w}'/>"
        f"<Cell N='Height' V='{h}'/>"
        f"<Cell N='LocPinX' V='{w/2}'/>"
        f"<Cell N='LocPinY' V='{h/2}'/>"
        f"<Cell N='LayerMember' V='{layer}'/>"
        f"<Cell N='LineWeight' V='0.006944444444444444' U='PT' F='Inh'/>"
        f"<Cell N='LineColor' V='#ffffff' F='Inh'/>"
        f"<Cell N='FillBkgnd' V='#4d4d4d' F='Inh'/>"
        f"{extra}"
        f"<Section N='Character'><Row IX='0'><Cell N='Color' V='#ffffff' F='Inh'/>"
        f"<Cell N='Size' V='0.1111111111111111' U='PT'/></Row></Section>"
        f"<Text><cp IX='0'/>{txt}</Text></Shape>"
    )


def connector_xml(sid: int, src_id: int, dst_id: int, a: Node, b: Node, edge: Edge) -> str:
    bx, by, ex, ey = a.x, a.y - a.h / 2, b.x, b.y + b.h / 2
    # horizontal if same row
    if abs(a.y - b.y) < 0.15:
        bx, by = a.x + a.w / 2, a.y
        ex, ey = b.x - b.w / 2, b.y
    width = ex - bx if abs(ex - bx) > 0.01 else 0.1968503937007874
    height = ey - by
    pinx = (bx + ex) / 2
    piny = (by + ey) / 2
    locx = width / 2
    locy = height / 2
    label = escape(edge.label)
    text = f"<Text>{label}</Text>" if edge.label else "<Text/>"
    return (
        f"<Shape ID='{sid}' NameU='Dynamic connector.{sid}' Name='Conector dinâmico.{sid}' Type='Shape' Master='{M_CONN}'>"
        f"<Cell N='PinX' V='{pinx}' F='Inh'/>"
        f"<Cell N='PinY' V='{piny}' F='Inh'/>"
        f"<Cell N='Width' V='{width}' F='GUARD(EndX-BeginX)'/>"
        f"<Cell N='Height' V='{height}' F='GUARD(EndY-BeginY)'/>"
        f"<Cell N='LocPinX' V='{locx}' F='Inh'/>"
        f"<Cell N='LocPinY' V='{locy}' F='Inh'/>"
        f"<Cell N='BeginX' V='{bx}' F='_WALKGLUE(BegTrigger,EndTrigger,WalkPreference)'/>"
        f"<Cell N='BeginY' V='{by}' F='_WALKGLUE(BegTrigger,EndTrigger,WalkPreference)'/>"
        f"<Cell N='EndX' V='{ex}' F='_WALKGLUE(EndTrigger,BegTrigger,WalkPreference)'/>"
        f"<Cell N='EndY' V='{ey}' F='_WALKGLUE(EndTrigger,BegTrigger,WalkPreference)'/>"
        f"<Cell N='LayerMember' V='1'/>"
        f"<Cell N='WalkPreference' V='3'/>"
        f"<Cell N='BegTrigger' V='2' F='_XFTRIGGER(Sheet.{src_id}!EventXFMod)'/>"
        f"<Cell N='EndTrigger' V='2' F='_XFTRIGGER(Sheet.{dst_id}!EventXFMod)'/>"
        f"<Cell N='LineWeight' V='0.006944444444444444' U='PT' F='Inh'/>"
        f"<Cell N='LineColor' V='{edge.color}' F='Inh'/>"
        f"<Cell N='EndArrow' V='4' F='Inh'/>"
        f"<Cell N='ConFixedCode' V='6'/>"
        f"<Section N='Character'><Row IX='0'><Cell N='Color' V='#1176a7'/><Cell N='Size' V='0.09722222222222222' U='PT'/></Row></Section>"
        f"<Section N='Geometry' IX='0'><Row T='MoveTo' IX='1'/><Row T='LineTo' IX='2'><Cell N='X' V='{width}'/><Cell N='Y' V='{height}'/></Row></Section>"
        f"{text}</Shape>"
    )


def gradient_xml(page_w: float, page_h: float) -> str:
    return (
        f"<Shape ID='1' NameU='Center Gradient' Name='Gradiente Central' Type='Shape' Master='{M_GRAD}'>"
        f"<Cell N='PinX' V='{page_w/2}'/>"
        f"<Cell N='PinY' V='{page_h/2}'/>"
        f"<Cell N='Width' V='{page_w}'/>"
        f"<Cell N='Height' V='{page_h}'/>"
        f"<Cell N='LocPinX' V='{page_w/2}'/>"
        f"<Cell N='LocPinY' V='{page_h/2}'/>"
        f"<Cell N='EventDrop' V='0'/>"
        f"<Cell N='EffectSchemeIndex' V='34'/>"
        f"<Cell N='FillForegnd' V='#79c8ee'/>"
        f"<Cell N='FillBkgnd' V='#d1ecf9'/>"
        f"<Section N='User'><Row N='msvVisioCreated'><Cell N='Value' V='0'/><Cell N='Prompt' V=''/></Row></Section>"
        f"</Shape>"
    )


def page_contents(page: Page) -> tuple[str, dict[str, int]]:
    ids: dict[str, int] = {}
    sid = 2
    parts = [gradient_xml(page.width, page.height)]
    for n in page.nodes:
        ids[n.key] = sid
        parts.append(shape_xml(sid, n))
        sid += 1
    connects = []
    for e in page.edges:
        if e.src not in ids or e.dst not in ids:
            continue
        src, dst = ids[e.src], ids[e.dst]
        a = next(x for x in page.nodes if x.key == e.src)
        b = next(x for x in page.nodes if x.key == e.dst)
        cid = sid
        parts.append(connector_xml(cid, src, dst, a, b, e))
        connects.append(
            f"<Connect FromSheet='{cid}' FromCell='BeginX' FromPart='9' ToSheet='{src}' ToCell='PinY' ToPart='3'/>"
            f"<Connect FromSheet='{cid}' FromCell='EndX' FromPart='12' ToSheet='{dst}' ToCell='PinY' ToPart='3'/>"
        )
        sid += 1
    xml = (
        "<?xml version='1.0' encoding='utf-8' ?>"
        f"<PageContents xmlns='{NS}' xmlns:r='{REL_NS}' xml:space='preserve'>"
        f"<Shapes>{''.join(parts)}</Shapes>"
        f"<Connects>{''.join(connects)}</Connects>"
        "</PageContents>"
    )
    return xml, ids


def page_sheet(page: Page, pid: int, unique: str) -> str:
    return (
        f"<Page ID='{pid}' NameU='{escape(page.name_u)}' IsCustomNameU='1' Name='{escape(page.name)}' IsCustomName='1' "
        f"ViewScale='0.65' ViewCenterX='{page.width/2}' ViewCenterY='{page.height/2}'>"
        f"<PageSheet LineStyle='0' FillStyle='0' TextStyle='0' UniqueID='{unique}'>"
        f"<Cell N='PageWidth' V='{page.width}'/>"
        f"<Cell N='PageHeight' V='{page.height}'/>"
        f"<Cell N='ShdwOffsetX' V='0.1181102362204724'/>"
        f"<Cell N='ShdwOffsetY' V='-0.1181102362204724'/>"
        f"<Cell N='PageScale' V='1' U='IN_F'/>"
        f"<Cell N='DrawingScale' V='1' U='IN_F'/>"
        f"<Cell N='DrawingSizeType' V='0'/>"
        f"<Cell N='DrawingScaleType' V='0'/>"
        f"<Cell N='InhibitSnap' V='0'/>"
        f"<Cell N='UIVisibility' V='0'/>"
        f"<Cell N='PrintPageOrientation' V='2'/>"
        f"<Cell N='PlaceStyle' V='2'/>"
        f"<Cell N='RouteStyle' V='6'/>"
        f"<Cell N='ColorSchemeIndex' V='34'/>"
        f"<Cell N='EffectSchemeIndex' V='34'/>"
        f"<Cell N='ThemeIndex' V='34'/>"
        f"<Section N='Layer'>"
        f"<Row IX='0'><Cell N='Name' V='Fluxograma'/><Cell N='Visible' V='1'/><Cell N='Print' V='1'/><Cell N='NameUniv' V='Flowchart'/></Row>"
        f"<Row IX='1'><Cell N='Name' V='Conector'/><Cell N='Visible' V='1'/><Cell N='Print' V='1'/><Cell N='NameUniv' V='Connector'/></Row>"
        f"</Section></PageSheet><Rel r:id='rId{pid+1}'/></Page>"
    )


def page_rels() -> str:
    rels = [
        f'<Relationship Id="rId{i}" Type="http://schemas.microsoft.com/visio/2010/relationships/master" Target="../masters/{fn}"/>'
        for i, fn in MASTER_RELS
    ]
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels)
        + "</Relationships>"
    )


def col(x0: float, y_top: float, gap: float, items: list[tuple]) -> tuple[list[Node], list[Edge]]:
    """items: (key, text, kind, href?) sequential top-down."""
    nodes, edges = [], []
    y = y_top
    prev = None
    for it in items:
        key, text, kind = it[0], it[1], it[2]
        href = it[3] if len(it) > 3 else None
        h = 0.55 if kind in ("start", "end") else 0.68
        nodes.append(Node(key, text, kind, x0, y, 2.15, h, href))
        if prev:
            edges.append(Edge(prev, key))
        prev = key
        y -= gap
    return nodes, edges

def q(key: str, method: str, path: str, kind: str, x: float, y: float, extra: str = "", w: float = 2.55, h: float = 0.78, href: str | None = None) -> Node:
    line2 = f"\n{extra}" if extra else ""
    return Node(key, f"{method} {path}{line2}", kind, x, y, w, h, href)


def build_pages() -> list[Page]:
    pages: list[Page] = []

    # ---- Legenda ----
    pages.append(
        Page(
            "Legenda",
            "Legenda",
            13.2,
            8.6,
            [
                Node("l0", "Início / Fim", "start", 2.2, 7.5, 2.3, 0.5),
                Node("l1", "Processo (BFF / FastAPI)", "process", 5.1, 7.5, 2.6, 0.55),
                Node("l2", "Subprocesso (outra página)", "sub", 8.1, 7.5, 2.6, 0.55),
                Node("l3", "Decisão", "decision", 11.0, 7.5, 2.2, 0.7),
                Node("l4", "SQL / Redis interno", "data", 2.2, 6.3, 2.6, 0.65),
                Node("l5", "API / banco externo", "ext", 5.4, 6.3, 2.6, 0.65),
                Node("l6", "Onboarding 2026-09-09\nRender FORA · Vercel no ar", "note", 9.2, 6.3, 4.0, 0.7),
            ],
            [],
        )
    )

    # ---- Foto produção ----
    n, e = col(
        3.2,
        12.4,
        0.95,
        [
            ("f0", "Browser", "start"),
            ("f1", "GET https://judgetcg.com.br/\nCloudflare + Vercel HTML 200", "process"),
            ("f2", "Homepage prerender\nx-nextjs-prerender: 1", "process"),
            ("f3", "BFF Next /api/*", "process"),
        ],
    )
    n += [
        Node("f4", "GET /api/health", "sub", 6.6, 9.55, 2.4, 0.7, "Health BFF"),
        Node("f5", "GET /api/games", "sub", 6.6, 8.4, 2.4, 0.65, "Games"),
        Node("f6", "POST /api/judge/query", "sub", 6.6, 7.25, 2.5, 0.7, "Judge BFF"),
        Node("f7", "GET /api/proxy/:path*\nrewrite → Render", "sub", 6.6, 6.05, 2.6, 0.75, "Proxy Render"),
        Node("f8", "API_PROXY_TARGET\nseekguidance.onrender.com", "ext", 9.8, 6.05, 2.7, 0.8),
        Node("f9", "Render no-server / 404\nou Service Suspended", "decision", 9.8, 4.7, 2.7, 0.85),
        Node("f10", "LPC = 0 · Beta not started", "note", 3.2, 4.7, 2.6, 0.7),
        Node("f11", "Fim (site no ar,\nAPI morta)", "end", 3.2, 3.4, 2.4, 0.6),
    ]
    e += [
        Edge("f3", "f4"),
        Edge("f3", "f5"),
        Edge("f3", "f6"),
        Edge("f3", "f7"),
        Edge("f7", "f8"),
        Edge("f8", "f9"),
        Edge("f9", "f11", "404/503", "#e51400"),
        Edge("f4", "f11"),
        Edge("f5", "f11"),
        Edge("f6", "f11"),
    ]
    pages.append(Page("Foto produção", "Foto produção", 13.4, 13.6, n, e))

    # ---- Health BFF (todas as consultas do health) ----
    n = [
        Node("h0", "GET /api/health", "start", 3.2, 12.2, 2.4, 0.5),
        Node("h1", "health/route.ts\nPromise.all 3 checks", "process", 3.2, 11.1, 2.5, 0.7),
        Node("h2", "checkSupabaseCatalog", "process", 3.2, 9.9, 2.5, 0.65),
        Node("h3", "SQL PostgREST\nSELECT id FROM tcg_judge.card_catalog LIMIT 1", "data", 6.6, 9.9, 3.2, 0.8),
        Node("h4", "card_catalog ok?", "decision", 3.2, 8.6, 2.3, 0.7),
        Node("h5", "checkApiHealth\nGET {API_PROXY_TARGET}/v1/health", "ext", 3.2, 7.3, 2.7, 0.8),
        Node("h6", "HTTP 200 FastAPI?", "decision", 3.2, 6.0, 2.3, 0.7),
        Node("h7", "checkRedis\nUpstash REST PING", "ext", 7.0, 7.3, 2.5, 0.75),
        Node("h8", "PING ok?", "decision", 7.0, 6.0, 2.2, 0.65),
        Node("h9", "database: error\n(observado 2026-09-09)", "process", 6.6, 8.6, 2.6, 0.7),
        Node("h10", "catalog_api: HTTP 404\n(Render no-server)", "process", 3.2, 4.7, 2.6, 0.75),
        Node("h11", "redis: ok", "process", 7.0, 4.7, 2.2, 0.55),
        Node("h12", "HTTP 503 degraded\nAPI+DB offline · Redis online", "end", 5.0, 3.4, 3.2, 0.7),
    ]
    e = [
        Edge("h0", "h1"),
        Edge("h1", "h2"),
        Edge("h2", "h3", "SQL"),
        Edge("h3", "h4"),
        Edge("h4", "h9", "Não", "#e51400"),
        Edge("h1", "h5"),
        Edge("h5", "h6"),
        Edge("h6", "h10", "Não 404", "#e51400"),
        Edge("h1", "h7"),
        Edge("h7", "h8"),
        Edge("h8", "h11", "Sim", "#75d175"),
        Edge("h9", "h12"),
        Edge("h10", "h12"),
        Edge("h11", "h12"),
    ]
    pages.append(Page("Health BFF", "Health BFF", 12.6, 13.4, n, e))

    # ---- Proxy Render ----
    n, e = col(
        3.4,
        10.4,
        1.0,
        [
            ("r0", "Cliente ou BFF", "start"),
            ("r1", "GET /api/proxy/:path*\nnext.config.mjs rewrite", "process"),
            ("r2", "Destino API_PROXY_TARGET\n+ /:path*", "ext"),
            ("r3", "seekguidance.onrender.com\nx-render-routing: no-server", "ext"),
            ("r4", "Serviço existe?", "decision"),
        ],
    )
    n += [
        Node("r5", "404 Not Found\nGET /api/proxy/v1/health", "end", 3.4, 4.4, 2.6, 0.7),
        Node("r6", "seekguidance2-66dz\n503 Service Suspended", "end", 7.0, 6.4, 2.6, 0.75),
        Node("r7", "tcg-judge-api.onrender.com\n404", "end", 7.0, 5.1, 2.6, 0.65),
    ]
    e += [
        Edge("r4", "r5", "Não", "#e51400"),
        Edge("r3", "r6"),
        Edge("r3", "r7"),
    ]
    pages.append(Page("Proxy Render", "Proxy Render", 11.2, 11.6, n, e))

    # ---- Games ----
    n, e = col(
        3.3,
        11.0,
        0.95,
        [
            ("g0", "GET /api/games", "start"),
            ("g1", "BFF api/games/route.ts", "process"),
            ("g2", "GET {API}/runtime/judge/catalog/games", "ext"),
            ("g3", "res.ok?", "decision"),
        ],
    )
    n += [
        Node("g4", "GET {API}/runtime/judge/games", "ext", 6.8, 8.15, 2.7, 0.7),
        Node("g5", "JSON games[]?", "decision", 3.3, 6.2, 2.3, 0.7),
        Node("g6", "catch → HTTP 503\n{ games: [] }", "end", 3.3, 5.0, 2.4, 0.7),
        Node("g7", "200 lista de jogos\n(só com FastAPI up)", "end", 6.8, 6.2, 2.5, 0.7),
    ]
    e += [
        Edge("g3", "g5", "Sim", "#75d175"),
        Edge("g3", "g4", "Não"),
        Edge("g4", "g5"),
        Edge("g5", "g7", "Sim", "#75d175"),
        Edge("g5", "g6", "Não", "#e51400"),
    ]
    pages.append(Page("Games", "Games", 11.2, 12.2, n, e))

    # ---- Judge BFF (consultas até o gate) ----
    n, e = col(
        3.2,
        12.6,
        0.92,
        [
            ("j0", "POST /api/judge/query", "start"),
            ("j1", "Body JSON tem tcg?", "decision"),
        ],
    )
    n += [
        Node("j2", "400 Campo tcg é obrigatório", "end", 6.8, 11.68, 2.6, 0.6),
        Node("j3", "getAuthenticatedUserId\nSupabase SSR cookie", "ext", 3.2, 10.6, 2.6, 0.75),
        Node("j4", "resolveRulesAccess(userId)", "process", 3.2, 9.5, 2.5, 0.65),
        Node("j5", "rules.allowed?", "decision", 3.2, 8.4, 2.3, 0.7),
        Node("j6", "403 login_required\n(observado sem sessão)", "end", 6.8, 8.4, 2.6, 0.75),
        Node("j7", "checkAndReserveQuestion\nplano / cookie diário", "data", 3.2, 7.2, 2.6, 0.7),
        Node("j8", "POST {API_BASE}/runtime/judge/query", "sub", 3.2, 6.0, 2.7, 0.75, "Judge RAG"),
        Node("j9", "Upstream Render morto\nfalha após login", "end", 6.8, 6.0, 2.5, 0.7),
    ]
    e += [
        Edge("j1", "j2", "Não", "#e51400"),
        Edge("j1", "j3", "Sim", "#75d175"),
        Edge("j3", "j4"),
        Edge("j4", "j5"),
        Edge("j5", "j6", "Não", "#e51400"),
        Edge("j5", "j7", "Sim", "#75d175"),
        Edge("j7", "j8"),
        Edge("j8", "j9", "prod 2026-09-09", "#e51400"),
    ]
    pages.append(Page("Judge BFF", "Judge BFF", 11.6, 13.8, n, e))

    # ---- Judge RAG consultas (local / quando API up) ----
    n = [
        Node("s0", "POST /runtime/judge/query", "start", 2.5, 16.2, 2.4, 0.5),
        Node("s1", "Middlewares FastAPI", "process", 2.5, 15.1, 2.3, 0.6),
        Node("s2", "JWT fail-closed (prod)\nX-Judge-User-Id RBAC", "decision", 2.5, 13.9, 2.4, 0.75),
        Node("s3", "Rate limit Redis\nREDIS_URL", "data", 5.6, 13.9, 2.3, 0.7),
        Node("s4", "GameConfiguration\nTCG válido?", "decision", 2.5, 12.6, 2.4, 0.75),
        Node("s5", "Cache semântico Redis HIT?", "decision", 2.5, 11.3, 2.5, 0.75),
        Node("s6", "Resposta cache", "process", 5.6, 11.3, 2.2, 0.6),
        Node("s7", "RagOrchestrator", "process", 2.5, 10.0, 2.3, 0.6),
        Node("s8", "HybridRetriever", "process", 2.5, 8.8, 2.3, 0.6),
        Node("s9", "SQL pgvector\nembedding <=> vector", "data", 5.6, 8.8, 2.5, 0.75),
        Node("s10", "SQL FTS\nto_tsvector / plainto_tsquery", "data", 8.6, 8.8, 2.6, 0.75),
        Node("s11", "Há hits?", "decision", 2.5, 7.5, 2.2, 0.7),
        Node("s12", "LLM OpenAI\nchat.completions", "ext", 2.5, 6.2, 2.4, 0.7),
        Node("s13", "Sem OPENAI_API_KEY\nstub confiança baixa", "note", 5.6, 6.2, 2.5, 0.75),
        Node("s14", "INSERT feedback / cache", "data", 2.5, 4.9, 2.4, 0.65),
        Node("s15", "JSON JudgeQueryResponse", "end", 2.5, 3.7, 2.4, 0.6),
    ]
    e = [
        Edge("s0", "s1"),
        Edge("s1", "s2"),
        Edge("s2", "s3"),
        Edge("s2", "s4", "OK", "#75d175"),
        Edge("s4", "s5", "Sim", "#75d175"),
        Edge("s5", "s6", "HIT", "#75d175"),
        Edge("s5", "s7", "MISS"),
        Edge("s6", "s15"),
        Edge("s7", "s8"),
        Edge("s8", "s9", "SQL"),
        Edge("s8", "s10", "SQL"),
        Edge("s9", "s11"),
        Edge("s10", "s11"),
        Edge("s11", "s12", "Sim", "#75d175"),
        Edge("s12", "s13"),
        Edge("s12", "s14"),
        Edge("s14", "s15"),
    ]
    pages.append(Page("Judge RAG", "Judge RAG", 12.2, 17.4, n, e))

    # ---- Dois bancos ----
    n = [
        Node("b0", "Qual banco?", "decision", 6.4, 10.4, 2.4, 0.75),
        Node("b1", "Postgres tcg_judge\ninfra/db/*.sql", "data", 3.0, 8.8, 2.6, 0.8),
        Node("b2", "Supabase public+auth\nmigrations RLS", "data", 9.6, 8.8, 2.6, 0.8),
        Node("b3", "SELECT games / documents / chunks\npgvector FTS", "data", 3.0, 7.3, 2.7, 0.8),
        Node("b4", "listings · stores · companies\ncard_catalog PostgREST", "data", 9.6, 7.3, 2.7, 0.8),
        Node("b5", "Local: Docker/cluster\nDATABASE_URL asyncpg", "process", 3.0, 5.8, 2.6, 0.75),
        Node("b6", "Prod health: erro\ncard_catalog inacessível", "process", 9.6, 5.8, 2.7, 0.75),
        Node("b7", "API Render FORA\nconsultas Judge indisponíveis", "end", 3.0, 4.4, 2.7, 0.7),
        Node("b8", "Auth UI pode viver\nsem FastAPI", "end", 9.6, 4.4, 2.5, 0.65),
    ]
    e = [
        Edge("b0", "b1", "Judge / RAG"),
        Edge("b0", "b2", "Marketplace / Auth"),
        Edge("b1", "b3"),
        Edge("b2", "b4"),
        Edge("b3", "b5"),
        Edge("b4", "b6"),
        Edge("b5", "b7"),
        Edge("b6", "b8"),
    ]
    pages.append(Page("Dois bancos", "Dois bancos", 13.2, 11.8, n, e))

    # ---- Auth ----
    n, e = col(
        3.3,
        10.8,
        0.95,
        [
            ("a0", "GET /entrar", "start"),
            ("a1", "Supabase Auth GoTrue\nNEXT_PUBLIC_SUPABASE_*", "ext"),
            ("a2", "Cookie @supabase/ssr", "process"),
            ("a3", "Sessão?", "decision"),
        ],
    )
    n += [
        Node("a4", "Gate /judge /decks /seller", "end", 6.8, 7.95, 2.5, 0.65),
        Node("a5", "BFF Authorization Bearer\n+ X-Judge-User-Id", "process", 3.3, 6.0, 2.6, 0.75),
        Node("a6", "FastAPI JWT fail-closed\n(ENVIRONMENT=production)", "process", 3.3, 4.8, 2.7, 0.75),
        Node("a7", "Fim autenticado", "end", 3.3, 3.6, 2.2, 0.5),
    ]
    e += [
        Edge("a3", "a4", "Não", "#e51400"),
        Edge("a3", "a5", "Sim", "#75d175"),
        Edge("a5", "a6"),
        Edge("a6", "a7"),
    ]
    pages.append(Page("Auth", "Auth", 11.0, 12.0, n, e))

    # ---- Marketplace consultas ----
    n = [
        Node("m0", "Buyer home / PDP RSC", "start", 3.0, 11.2, 2.4, 0.55),
        Node("m1", "BFF listings / catalog cards", "process", 3.0, 10.1, 2.5, 0.65),
        Node("m2", "SQL public + RLS\nlistings stores", "data", 3.0, 8.9, 2.5, 0.7),
        Node("m3", "Seller CNPJ?\nidentity_platform", "decision", 3.0, 7.6, 2.4, 0.75),
        Node("m4", "SQL companies\nvalidação local dígitos", "data", 6.2, 7.6, 2.6, 0.75),
        Node("m5", "GET /api/geo/cep/[cep]\nBrasilAPI /api/cep/v2", "ext", 9.4, 7.6, 2.6, 0.8),
        Node("m6", "Stripe Checkout / Connect\nPIX webhook", "ext", 3.0, 6.2, 2.5, 0.75),
        Node("m7", "Melhor Envio\nfreight_quote.py", "ext", 6.2, 6.2, 2.4, 0.7),
        Node("m8", "Evento LPC\nsem intervenção humana", "process", 3.0, 4.9, 2.5, 0.7),
        Node("m9", "Prod: FastAPI fora\ncheckout morto", "end", 6.2, 4.9, 2.4, 0.65),
        Node("m10", "LPC hoje = 0", "end", 3.0, 3.6, 2.2, 0.5),
    ]
    e = [
        Edge("m0", "m1"),
        Edge("m1", "m2", "SQL"),
        Edge("m2", "m3"),
        Edge("m3", "m4", "credenciar"),
        Edge("m0", "m5", "CEP"),
        Edge("m3", "m6", "buyer paga"),
        Edge("m6", "m7"),
        Edge("m6", "m8"),
        Edge("m6", "m9", "prod", "#e51400"),
        Edge("m8", "m10"),
    ]
    pages.append(Page("Marketplace", "Marketplace", 13.0, 12.4, n, e))

    # ---- Catálogo consultas ----
    n = [
        Node("c0", "Job / request catálogo", "start", 2.6, 10.0, 2.3, 0.5),
        Node("c1", "Product Catalog Sync\nPRODUCT_CATALOG_DATABASE_URL", "process", 2.6, 8.8, 2.6, 0.75),
        Node("c2", "Tipo?", "decision", 2.6, 7.5, 2.1, 0.65),
        Node("c3", "Cartas\nScryfall pokemontcg.io\nYGOPRODeck lorcast OPTCG", "ext", 5.6, 7.5, 2.7, 0.95),
        Node("c4", "Selados\nTCGCSV → TCGplayer CDN", "ext", 8.7, 7.5, 2.5, 0.8),
        Node("c5", "Acessórios\nShopify Woo Shopware Tray", "ext", 11.6, 7.5, 2.6, 0.85),
        Node("c6", "SQL product_catalog\n+ media.assets", "data", 2.6, 5.7, 2.5, 0.7),
        Node("c7", "Meilisearch (opcional)", "ext", 5.6, 5.7, 2.4, 0.65),
        Node("c8", "Cron catalog-health-ping\nGET /runtime/judge/catalog/health?lite=1", "ext", 8.8, 5.7, 3.0, 0.8),
        Node("c9", "Prod: job nunca verde\nRender + secret ausentes", "end", 2.6, 4.3, 2.6, 0.7),
    ]
    e = [
        Edge("c0", "c1"),
        Edge("c1", "c2"),
        Edge("c2", "c3", "single"),
        Edge("c2", "c4", "sealed"),
        Edge("c2", "c5", "accessory"),
        Edge("c3", "c6"),
        Edge("c4", "c6"),
        Edge("c5", "c6"),
        Edge("c6", "c7"),
        Edge("c1", "c8", "Actions */14"),
        Edge("c8", "c9", "404 Render", "#e51400"),
        Edge("c6", "c9"),
    ]
    pages.append(Page("Catálogo", "Catálogo", 15.0, 11.2, n, e))

    # ---- Local (consultas que funcionam) ----
    n, e = col(
        3.3,
        12.4,
        0.88,
        [
            ("l0", "Dev local (caminho oficial)", "start"),
            ("l1", "Postgres 16 + pgvector\npsql seed infra/db/*.sql", "data"),
            ("l2", "Redis 6379 ou compose 6380\nREDIS_URL alinhado", "data"),
            ("l3", "uvicorn :8000\nGET /health GET /v1/health", "process"),
            ("l4", "GET /v1/games\n14 slugs do seed", "data"),
            ("l5", "GET /docs OpenAPI", "process"),
            ("l6", "Next :3000\nAPI_PROXY_TARGET=http://127.0.0.1:8000", "process"),
            ("l7", "Rewrite /api/proxy/* → :8000", "process"),
            ("l8", "Judge stub se sem OPENAI_API_KEY", "end"),
        ],
    )
    pages.append(Page("Stack local", "Stack local", 10.4, 13.6, n, e))

    # ---- Crons / orçamento ----
    n = [
        Node("o0", "GitHub Actions crons", "start", 3.0, 10.6, 2.4, 0.5),
        Node("o1", "catalog-health-ping */14\nGET …/catalog/health?lite=1", "ext", 3.0, 9.3, 2.7, 0.8),
        Node("o2", "expire-checkouts */5\nPOST …/checkout/expire-stale", "ext", 6.2, 9.3, 2.7, 0.8),
        Node("o3", "catalog-sync\nPRODUCT_CATALOG_DATABASE_URL", "data", 9.4, 9.3, 2.7, 0.8),
        Node("o4", "Minutos Actions esgotados\nsteps: [] sem logs", "decision", 3.0, 7.7, 2.6, 0.8),
        Node("o5", "expire: 401 sem JWT\n(prefixo mutação)", "process", 6.2, 7.7, 2.6, 0.75),
        Node("o6", "Vercel Hobby 10s\nAPI_FETCH_TIMEOUT_MS=7000", "process", 9.4, 7.7, 2.6, 0.8),
        Node("o7", "Não religar crons\npara mascarar 503", "end", 6.2, 6.1, 2.6, 0.7),
    ]
    e = [
        Edge("o0", "o1"),
        Edge("o0", "o2"),
        Edge("o0", "o3"),
        Edge("o1", "o4"),
        Edge("o2", "o5"),
        Edge("o3", "o4"),
        Edge("o6", "o7"),
        Edge("o4", "o7"),
        Edge("o5", "o7"),
    ]
    pages.append(Page("Crons orçamento", "Crons orçamento", 13.2, 11.8, n, e))

    # ---- Inventário (mapa de todas as consultas) ----
    n = [
        Node("i0", "Inventário de consultas\n(transpor no Visio via Excel)", "note", 7.2, 11.6, 4.2, 0.7),
        Node("i1", "GET /  ·  GET /api/health", "process", 2.4, 10.2, 2.6, 0.7),
        Node("i2", "GET /api/games", "process", 5.4, 10.2, 2.3, 0.55),
        Node("i3", "POST /api/judge/query", "process", 8.2, 10.2, 2.5, 0.55),
        Node("i4", "GET /api/proxy/*", "process", 11.2, 10.2, 2.4, 0.55),
        Node("i5", "GET /v1/health  /health", "ext", 2.4, 8.8, 2.6, 0.65),
        Node("i6", "GET /v1/games  /docs", "ext", 5.4, 8.8, 2.4, 0.65),
        Node("i7", "POST /runtime/judge/query", "ext", 8.2, 8.8, 2.6, 0.65),
        Node("i8", "GET catalog/games · games", "ext", 11.2, 8.8, 2.5, 0.7),
        Node("i9", "SQL card_catalog LIMIT 1", "data", 2.4, 7.4, 2.6, 0.65),
        Node("i10", "SQL chunks <=> vector", "data", 5.4, 7.4, 2.4, 0.65),
        Node("i11", "SQL to_tsvector FTS", "data", 8.2, 7.4, 2.4, 0.65),
        Node("i12", "SQL companies / listings", "data", 11.2, 7.4, 2.5, 0.65),
        Node("i13", "Upstash PING", "ext", 2.4, 6.0, 2.4, 0.55),
        Node("i14", "Redis cache / rate limit", "data", 5.4, 6.0, 2.4, 0.65),
        Node("i15", "OpenAI embeddings+chat", "ext", 8.2, 6.0, 2.5, 0.65),
        Node("i16", "Supabase GoTrue", "ext", 11.2, 6.0, 2.4, 0.55),
        Node("i17", "BrasilAPI CEP", "ext", 2.4, 4.6, 2.4, 0.55),
        Node("i18", "Stripe / PIX", "ext", 5.4, 4.6, 2.3, 0.55),
        Node("i19", "Melhor Envio", "ext", 8.2, 4.6, 2.3, 0.55),
        Node("i20", "Scryfall / TCGCSV / Shopify", "ext", 11.2, 4.6, 2.6, 0.7),
        Node("i21", "GET catalog/health?lite=1", "ext", 2.4, 3.2, 2.6, 0.65),
        Node("i22", "POST checkout/expire-stale", "ext", 5.4, 3.2, 2.6, 0.65),
        Node("i23", "Meilisearch", "ext", 8.2, 3.2, 2.3, 0.5),
        Node("i24", "Planilha Catalogo_Consultas.xlsx", "note", 11.2, 3.2, 2.6, 0.7),
    ]
    pages.append(Page("Inventário", "Inventário", 14.6, 12.8, n, []))

    return pages


# Linhas para Visio Data Visualizer / transposição manual
# Shape Type: Start/End, Process, Decision, Data, External Data, Subprocess
QUERY_FLOWS: dict[str, list[tuple[str, str, str, str, str, str]]] = {}


def _flow(name: str, rows: list[tuple[str, str, str, str, str, str]]) -> None:
    QUERY_FLOWS[name] = rows


_flow(
    "Q01_Health",
    [
        ("H01", "GET /api/health", "H02", "HTTP", "Start/End", "BFF"),
        ("H02", "Promise.all: Supabase + FastAPI + Upstash", "H03;H06;H09", "paralelo", "Process", "BFF"),
        ("H03", "PostgREST SELECT id FROM tcg_judge.card_catalog LIMIT 1", "H04", "SQL", "Data", "Supabase"),
        ("H04", "card_catalog ok?", "H05;H06", "Não;Sim", "Decision", "Supabase"),
        ("H05", "database: error (observado)", "H11", "", "Process", "Prod"),
        ("H06", "GET {API_PROXY_TARGET}/v1/health", "H07", "HTTP", "External Data", "FastAPI/Render"),
        ("H07", "FastAPI HTTP 200?", "H08;H11", "Sim;Não 404", "Decision", "Render"),
        ("H08", "catalog_api ok", "H11", "", "Process", "FastAPI"),
        ("H09", "Upstash Redis REST PING", "H10", "REST", "External Data", "Upstash"),
        ("H10", "PING ok? (observado Sim)", "H11", "Sim", "Decision", "Upstash"),
        ("H11", "HTTP 503 degraded JSON", "", "", "Start/End", "BFF"),
    ],
)
_flow(
    "Q02_Proxy",
    [
        ("P01", "GET /api/proxy/:path*", "P02", "HTTP", "Start/End", "Next rewrite"),
        ("P02", "Rewrite next.config.mjs → API_PROXY_TARGET/:path*", "P03", "", "Process", "Vercel"),
        ("P03", "GET https://seekguidance.onrender.com/v1/health", "P04", "HTTP", "External Data", "Render"),
        ("P04", "Serviço Render no ar?", "P05;P06", "Não;Sim", "Decision", "Render"),
        ("P05", "404 Not Found x-render-routing: no-server", "", "", "Start/End", "Prod 2026-09-09"),
        ("P06", "JSON health FastAPI", "", "", "Start/End", "Só se Render voltar"),
    ],
)
_flow(
    "Q03_Games",
    [
        ("G01", "GET /api/games", "G02", "HTTP", "Start/End", "BFF"),
        ("G02", "GET {API}/runtime/judge/catalog/games", "G03", "HTTP", "External Data", "FastAPI"),
        ("G03", "res.ok?", "G05;G04", "Sim;Não", "Decision", "BFF"),
        ("G04", "GET {API}/runtime/judge/games", "G05", "HTTP", "External Data", "FastAPI"),
        ("G05", "Há games[]?", "G06;G07", "Sim;Não/catch", "Decision", "BFF"),
        ("G06", "HTTP 200 { games }", "", "", "Start/End", "Local/API up"),
        ("G07", "HTTP 503 { games: [] }", "", "", "Start/End", "Prod observado"),
    ],
)
_flow(
    "Q04_Judge_BFF",
    [
        ("J01", "POST /api/judge/query", "J02", "HTTP", "Start/End", "BFF"),
        ("J02", "JSON.tcg presente?", "J03;J04", "Não;Sim", "Decision", "BFF"),
        ("J03", "400 Campo tcg é obrigatório", "", "", "Start/End", "Prod observado"),
        ("J04", "getAuthenticatedUserId (Supabase cookie)", "J05", "Auth", "External Data", "Supabase"),
        ("J05", "resolveRulesAccess", "J06", "", "Process", "BFF"),
        ("J06", "rules.allowed?", "J07;J08", "Não;Sim", "Decision", "BFF"),
        ("J07", "403 login_required", "", "", "Start/End", "Prod observado"),
        ("J08", "checkAndReserveQuestion (cookie plano)", "J09", "", "Data", "BFF"),
        ("J09", "POST {API_BASE}/runtime/judge/query", "J10", "HTTP", "Subprocess", "FastAPI"),
        ("J10", "Upstream disponível?", "J11;J12", "Não;Sim", "Decision", "Render"),
        ("J11", "Falha upstream (Render fora)", "", "", "Start/End", "Após login"),
        ("J12", "JSON JudgeQueryResponse", "", "", "Start/End", "Local/API up"),
    ],
)
_flow(
    "Q05_Judge_RAG",
    [
        ("S01", "POST /runtime/judge/query", "S02", "HTTP", "Start/End", "FastAPI"),
        ("S02", "Middlewares", "S03", "", "Process", "FastAPI"),
        ("S03", "JWT / rate limit OK?", "S04;S99", "OK;Bloqueada", "Decision", "Auth+Redis"),
        ("S04", "REDIS rate limit / revoke", "S05", "", "Data", "Redis"),
        ("S05", "GameConfiguration TCG válido?", "S06;S98", "Sim;Não", "Decision", "Judge"),
        ("S06", "Cache semântico HIT?", "S07;S08", "HIT;MISS", "Decision", "Redis"),
        ("S07", "Retornar cache", "S97", "", "Process", "Redis"),
        ("S08", "RagOrchestrator + HybridRetriever", "S09;S10", "SQL", "Process", "RAG"),
        ("S09", "SQL pgvector embedding <=> vector", "S11", "SQL", "Data", "tcg_judge.chunks"),
        ("S10", "SQL to_tsvector / plainto_tsquery", "S11", "SQL", "Data", "tcg_judge.chunks"),
        ("S11", "Há hits?", "S12;S98", "Sim;Não", "Decision", "Retrieval"),
        ("S12", "OpenAI chat.completions (ou stub)", "S13", "HTTPS", "External Data", "OpenAI"),
        ("S13", "Persistir feedback + cache", "S97", "SQL", "Data", "Postgres+Redis"),
        ("S97", "JudgeQueryResponse", "S99", "JSON", "Process", "API"),
        ("S98", "Erro / indisponível", "S99", "", "Process", "API"),
        ("S99", "Fim", "", "", "Start/End", "Usuário"),
    ],
)
_flow(
    "Q06_Local",
    [
        ("L01", "Início stack local", "L02", "", "Start/End", "Dev"),
        ("L02", "psql seed infra/db/init.sql … 06_*.sql", "L03", "SQL", "Data", "Postgres"),
        ("L03", "Redis PING (6379 ou 6380)", "L04", "", "Data", "Redis"),
        ("L04", "GET http://127.0.0.1:8000/health", "L05", "HTTP", "Process", "FastAPI"),
        ("L05", "GET /v1/health", "L06", "HTTP", "Process", "FastAPI"),
        ("L06", "GET /v1/games", "L07", "HTTP+SQL", "Data", "tcg_judge.games"),
        ("L07", "GET /docs", "L08", "HTTP", "Process", "OpenAPI"),
        ("L08", "Next API_PROXY_TARGET=http://127.0.0.1:8000", "L09", "", "Process", "BFF"),
        ("L09", "GET /api/proxy/v1/health → :8000/v1/health", "", "HTTP", "Subprocess", "Local"),
    ],
)
_flow(
    "Q07_Marketplace",
    [
        ("M01", "GET home / PDP RSC", "M02", "HTTP", "Start/End", "Next"),
        ("M02", "BFF listings / catalog cards", "M03", "HTTP", "Process", "BFF"),
        ("M03", "SQL listings stores RLS public", "M04", "SQL", "Data", "Supabase"),
        ("M04", "Seller CNPJ? companies", "M05;M06", "credenciar;comprar", "Decision", "identity_platform"),
        ("M05", "SQL INSERT/SELECT tcg_judge.companies (CNPJ local)", "M03", "SQL", "Data", "Postgres"),
        ("M06", "GET /api/geo/cep/{cep}", "M07", "HTTP", "Process", "BFF"),
        ("M07", "GET https://brasilapi.com.br/api/cep/v2/{cep}", "M08", "HTTPS", "External Data", "BrasilAPI"),
        ("M08", "Stripe Checkout / PIX / webhook", "M09", "HTTPS", "External Data", "Stripe"),
        ("M09", "Melhor Envio freight_quote", "M10", "HTTPS", "External Data", "Melhor Envio"),
        ("M10", "Pedido SQL + LPC", "", "", "Start/End", "Marketplace"),
    ],
)
_flow(
    "Q08_Catalogo",
    [
        ("C01", "Product Catalog Sync / request", "C02", "", "Start/End", "Worker"),
        ("C02", "Tipo de produto?", "C03;C04;C05", "carta;selado;acessório", "Decision", "Catalog"),
        ("C03", "GET Scryfall / pokemontcg.io / YGOPRODeck / lorcast / OPTCG", "C06", "HTTPS", "External Data", "APIs TCG"),
        ("C04", "GET TCGCSV → CDN TCGplayer", "C06", "HTTPS", "External Data", "TCGCSV"),
        ("C05", "GET Shopify / Woo / Shopware / Tray", "C06", "HTTPS", "External Data", "Fabricantes"),
        ("C06", "SQL product_catalog + media.assets", "C07", "SQL", "Data", "Catalog DB"),
        ("C07", "Meilisearch index (opcional)", "C08", "", "External Data", "Meilisearch"),
        ("C08", "GET /runtime/judge/catalog/health?lite=1 (cron */14)", "", "HTTP", "External Data", "Actions→Render"),
    ],
)
_flow(
    "Q09_Crons",
    [
        ("K01", "GitHub Actions schedule", "K02;K03;K04", "", "Start/End", "CI"),
        ("K02", "GET {SMOKE_API_URL}/runtime/judge/catalog/health?lite=1", "K05", "HTTP", "External Data", "Render"),
        ("K03", "POST {API}/runtime/judge/checkout/expire-stale", "K06", "HTTP", "External Data", "Render"),
        ("K04", "Catalog sync PRODUCT_CATALOG_DATABASE_URL", "K07", "SQL", "Data", "Catalog DB"),
        ("K05", "404 no-server (observado)", "", "", "Start/End", "Render"),
        ("K06", "401 sem JWT mesmo com API (prefixo mutação)", "", "", "Start/End", "Auth"),
        ("K07", "Falha sem secret / minutos Actions", "", "", "Start/End", "Orçamento"),
    ],
)
_flow(
    "Q10_Auth",
    [
        ("A01", "GET /entrar", "A02", "HTTP", "Start/End", "Next"),
        ("A02", "Supabase GoTrue Auth", "A03", "HTTPS", "External Data", "Supabase"),
        ("A03", "Cookie @supabase/ssr", "A04", "", "Process", "BFF"),
        ("A04", "Sessão válida?", "A05;A06", "Não;Sim", "Decision", "Auth"),
        ("A05", "Gate login /judge /decks /seller", "", "", "Start/End", "UI"),
        ("A06", "Authorization Bearer + X-Judge-User-Id", "A07", "", "Process", "BFF"),
        ("A07", "FastAPI JWT fail-closed (production)", "", "", "Start/End", "FastAPI"),
    ],
)


CATALOGO_CONSULTAS = [
    # metodo, path, destino, forma, prod, fonte onboarding
    ("GET", "https://judgetcg.com.br/", "Vercel+CF", "Process", "200 HTML prerender", "§1"),
    ("GET", "/api/health", "BFF Next", "Process", "503 degraded", "§1"),
    ("SQL", "tcg_judge.card_catalog SELECT id LIMIT 1", "Supabase PostgREST", "Data", "database error", "§1 health"),
    ("GET", "{API_PROXY_TARGET}/v1/health", "FastAPI/Render", "External Data", "HTTP 404", "§1 health"),
    ("REST", "Upstash Redis PING", "Upstash", "External Data", "ok", "§1 health"),
    ("GET", "/api/games", "BFF", "Process", "503 {games:[]}", "§1"),
    ("GET", "/runtime/judge/catalog/games", "FastAPI", "External Data", "404 via proxy", "§7.2"),
    ("GET", "/runtime/judge/games", "FastAPI", "External Data", "404 via proxy", "§7.2"),
    ("GET", "/api/proxy/:path*", "Rewrite→Render", "Process", "404 no-server", "§1 rewrite"),
    ("GET", "https://seekguidance.onrender.com/health", "Render", "External Data", "404 no-server", "§1"),
    ("GET", "https://seekguidance2-66dz.onrender.com/health", "Render", "External Data", "503 Suspended", "§1"),
    ("GET", "https://tcg-judge-api.onrender.com/v1/health", "Render legado", "External Data", "404", "§1"),
    ("POST", "/api/judge/query", "BFF", "Process", "400 sem tcg; 403 sem login", "§1"),
    ("AUTH", "getAuthenticatedUserId / GoTrue", "Supabase", "External Data", "gate login", "§9.4"),
    ("POST", "/runtime/judge/query", "FastAPI", "External Data", "só após login; upstream morto", "§9.1"),
    ("DATA", "Redis semantic cache + rate limit", "Redis API", "Data", "API fora em prod", "§9.1"),
    ("SQL", "chunks embedding <=> vector", "Postgres tcg_judge", "Data", "API fora", "§9.1"),
    ("SQL", "to_tsvector / plainto_tsquery", "Postgres tcg_judge", "Data", "API fora", "§9.1"),
    ("POST", "OpenAI chat.completions + embeddings", "OpenAI", "External Data", "sem chave = stub", "§8 §9.1"),
    ("GET", "http://127.0.0.1:8000/health", "FastAPI local", "Process", "quando stack sobe", "§10.2"),
    ("GET", "http://127.0.0.1:8000/v1/health", "FastAPI local", "Process", "quando stack sobe", "§10.2"),
    ("GET", "http://127.0.0.1:8000/v1/games", "FastAPI+SQL games", "Data", "14 slugs seed", "§10.2"),
    ("GET", "http://127.0.0.1:8000/docs", "OpenAPI", "Process", "local", "§10.2"),
    ("SQL", "psql infra/db/init.sql … 06_*.sql", "Postgres local", "Data", "seed", "§10.2"),
    ("GET", "/api/geo/cep/[cep]", "BFF", "Process", "não depende Render", "§2 §9.2"),
    ("GET", "https://brasilapi.com.br/api/cep/v2/{cep}", "BrasilAPI", "External Data", "externo", "§2"),
    ("SQL", "companies CNPJ (validação local)", "tcg_judge", "Data", "sem lookup Receita", "§2 §6"),
    ("SQL", "listings / stores RLS", "Supabase public", "Data", "vazio sem migrations", "§9.2"),
    ("HTTPS", "Stripe Checkout/Connect/PIX/webhook", "Stripe", "External Data", "morto sem API", "§9.2"),
    ("HTTPS", "Melhor Envio freight_quote", "Melhor Envio", "External Data", "morto sem API", "§9.2"),
    ("HTTPS", "Scryfall / pokemontcg.io / YGOPRODeck / lorcast / OPTCG", "APIs TCG", "External Data", "job sync", "§9.3"),
    ("HTTPS", "TCGCSV → TCGplayer CDN", "TCGCSV", "External Data", "selados", "§9.3"),
    ("HTTPS", "Shopify / Woo / Shopware / Tray", "Fabricantes", "External Data", "acessórios", "§9.3"),
    ("SQL", "product_catalog + media.assets", "Catalog DB", "Data", "precisa secret", "§9.3"),
    ("HTTP", "Meilisearch index/search", "Meilisearch", "External Data", "opcional / extra custo", "§8"),
    ("GET", "/runtime/judge/catalog/health?lite=1", "Render via Actions", "External Data", "cron */14 404", "§7.2 §8"),
    ("POST", "/runtime/judge/checkout/expire-stale", "Render via Actions", "External Data", "401 sem JWT", "§8"),
    ("SQL", "PRODUCT_CATALOG_DATABASE_URL sync", "Catalog DB", "Data", "nunca verde sem secret", "§8"),
]


def write_xlsx(dest: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()

    # Catalogo
    ws = wb.active
    ws.title = "Catalogo_Consultas"
    headers = [
        "Process Step ID",
        "Método",
        "Consulta (path / SQL / API)",
        "Destino",
        "Shape Type",
        "Estado prod 2026-09-09",
        "Fonte onboarding",
        "Function / Area",
    ]
    ws.append(headers)
    for i, row in enumerate(CATALOGO_CONSULTAS, 1):
        method, path, destine, shape, prod, fonte = row
        ws.append([f"Q{i:02d}", method, path, destine, shape, prod, fonte, destine])
    fill = PatternFill("solid", fgColor="1B4F72")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True)
    widths = [16, 12, 55, 22, 18, 36, 18, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = ws.max_row
    tab = Table(displayName="CatalogoConsultas", ref=f"A1:H{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.auto_filter.ref = f"A1:H{last}"
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "1B4F72"

    # Legenda formas Visio
    lg = wb.create_sheet("Legenda_Visio")
    lg.append(["Shape Type (Excel)", "Nome no stencil Judge_doc.vsdx", "Uso nas consultas"])
    for row in [
        ("Start/End", "Início/Término", "Início ou fim de um fluxo de consulta"),
        ("Process", "Processo", "BFF Next ou atividade FastAPI"),
        ("Subprocess", "Subprocesso", "Consulta detalhada em outra aba/página"),
        ("Decision", "Decisão", "HTTP ok? sessão? HIT cache? CNPJ?"),
        ("Data", "Dados", "SQL Postgres / Redis interno"),
        ("External Data", "Dados Externos", "Render, OpenAI, Stripe, BrasilAPI, Upstash, APIs TCG"),
    ]:
        lg.append(list(row))
    for cell in lg[1]:
        cell.fill = fill
        cell.font = font
    for i, w in enumerate([22, 36, 50], 1):
        lg.column_dimensions[get_column_letter(i)].width = w

    # Como importar
    how = wb.create_sheet("Como_importar_Visio")
    how.append(["Passo", "Ação no Microsoft Visio"])
    for row in [
        ("1", "Abrir JudgeTCG_Onboarding_Consultas.vsdx (já tem as páginas desenhadas)."),
        ("2", "Ou: Novo diagrama → Extensões → Visualizador de Dados → Criar → Fluxograma básico."),
        ("3", "Selecionar uma aba Q01_Health … Q10_Auth desta planilha como origem."),
        ("4", "Mapear colunas: Process Step ID, Process Step Description, Next Step ID, Connector Label, Shape Type, Function / Area."),
        ("5", "Shape Type deve bater com o stencil Fluxograma Básico (mesmos nomes em inglês da coluna)."),
        ("6", "Next Step ID aceita vários destinos separados por ponto-e-vírgula (ex.: H03;H06;H09)."),
        ("7", "A aba Catalogo_Consultas é o inventário completo para copiar/colar formas uma a uma se preferir."),
        ("8", "Fonte de verdade do conteúdo: docs/engineering/TEAM_ONBOARDING_ESTADO_ATUAL.md"),
    ]:
        how.append(list(row))
    for cell in how[1]:
        cell.fill = fill
        cell.font = font
    how.column_dimensions["A"].width = 8
    how.column_dimensions["B"].width = 110

    visio_headers = [
        "Process Step ID",
        "Process Step Description",
        "Next Step ID",
        "Connector Label",
        "Shape Type",
        "Function / Area",
    ]
    for sheet_name, rows in QUERY_FLOWS.items():
        sh = wb.create_sheet(sheet_name[:31])
        sh.append(visio_headers)
        for r in rows:
            sh.append(list(r))
        for cell in sh[1]:
            cell.fill = fill
            cell.font = font
        for i, w in enumerate([18, 62, 16, 18, 16, 22], 1):
            sh.column_dimensions[get_column_letter(i)].width = w
        end = sh.max_row
        t = Table(displayName=sheet_name.replace("-", "_")[:20], ref=f"A1:F{end}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        sh.add_table(t)
        sh.auto_filter.ref = f"A1:F{end}"
        sh.freeze_panes = "A2"

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def main() -> None:
    pages = build_pages()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    vsdx = OUT_DIR / "JudgeTCG_Onboarding_Consultas.vsdx"
    write_vsdx(pages, vsdx)
    (OUT_DIR / "JudgeTCG_Onboarding_Consultas.drawio").write_text(drawio_xml(pages), encoding="utf-8")
    html = html_preview(pages).replace(
        "JudgeTCG — Fluxograma funcional E2E",
        "JudgeTCG — Consultas E2E (onboarding)",
    ).replace(
        "Arquivo MS Visio: docs/architecture/visio/JudgeTCG_Fluxograma_E2E.vsdx",
        "MS Visio: docs/architecture/visio/JudgeTCG_Onboarding_Consultas.vsdx · Excel: JudgeTCG_Onboarding_Consultas.xlsx",
    )
    (OUT_DIR / "JudgeTCG_Onboarding_Consultas.html").write_text(html, encoding="utf-8")
    xlsx = OUT_DIR / "JudgeTCG_Onboarding_Consultas.xlsx"
    write_xlsx(xlsx)
    print(f"Wrote {vsdx} ({vsdx.stat().st_size} bytes)")
    print(f"Wrote {xlsx} ({xlsx.stat().st_size} bytes)")
    print(f"Pages: {[p.name for p in pages]}")
    print(f"Excel sheets: Catalogo_Consultas + {len(QUERY_FLOWS)} fluxos")
DRAW_COLORS = {
    "start": ("#e89ab8", "#5a1f35"),
    "end": ("#e89ab8", "#5a1f35"),
    "process": ("#7eb8da", "#123044"),
    "sub": ("#7dbe7d", "#143214"),
    "decision": ("#e89ab8", "#5a1f35"),
    "data": ("#e8a05c", "#3d2a12"),
    "ext": ("#d9894a", "#3d2a12"),
    "note": ("#f2f2f2", "#333333"),
}


def drawio_xml(pages: list[Page]) -> str:
    diagrams = []
    for pi, page in enumerate(pages):
        cells = [
            '<mxCell id="0"/>',
            '<mxCell id="1" parent="0"/>',
        ]
        key_id = {}
        cid = 2
        for n in page.nodes:
            key_id[n.key] = cid
            fill, font = DRAW_COLORS[n.kind]
            style = (
                f"rounded=1;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};"
                f"strokeColor=#ffffff;fontSize=11;align=center;"
            )
            if n.kind == "decision":
                style = f"rhombus;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};strokeColor=#ffffff;fontSize=11;"
            elif n.kind in ("start", "end"):
                style = f"ellipse;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};strokeColor=#ffffff;fontSize=11;"
            elif n.kind == "data":
                style = f"parallelogram;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};strokeColor=#ffffff;fontSize=11;"
            elif n.kind == "ext":
                style = f"cylinder3;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};strokeColor=#ffffff;fontSize=11;size=10;"
            elif n.kind == "sub":
                style = (
                    f"rounded=0;whiteSpace=wrap;html=1;fillColor={fill};fontColor={font};"
                    "strokeColor=#ffffff;fontSize=11;shape=partialRectangle;left=1;right=1;top=0;bottom=0;"
                )
            # Visio Y up → draw.io Y down
            x = int(n.x * 90)
            y = int((page.height - n.y) * 90)
            w = int(n.w * 90)
            h = int(n.h * 90)
            label = html.escape(n.text).replace("\n", "&#xa;")
            cells.append(
                f'<mxCell id="{cid}" value="{label}" style="{style}" vertex="1" parent="1">'
                f'<mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/></mxCell>'
            )
            cid += 1
        for e in page.edges:
            if e.src not in key_id or e.dst not in key_id:
                continue
            lbl = html.escape(e.label)
            cells.append(
                f'<mxCell id="{cid}" value="{lbl}" style="endArrow=block;html=1;strokeColor={e.color};fontSize=10;" '
                f'edge="1" parent="1" source="{key_id[e.src]}" target="{key_id[e.dst]}">'
                f'<mxGeometry relative="1" as="geometry"/></mxCell>'
            )
            cid += 1
        diagrams.append(
            f'<diagram id="p{pi}" name="{html.escape(page.name)}">'
            f'<mxGraphModel dx="1200" dy="800" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" '
            f'arrows="1" fold="1" page="1" pageScale="1" pageWidth="{int(page.width*96)}" pageHeight="{int(page.height*96)}" math="0" shadow="0">'
            f'<root>{"".join(cells)}</root></mxGraphModel></diagram>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<mxfile host="app.diagrams.net" type="device">'
        + "".join(diagrams)
        + "</mxfile>"
    )


def html_preview(pages: list[Page]) -> str:
    css = """
    body{font-family:Segoe UI,system-ui,sans-serif;background:#0f1720;color:#e8eef4;margin:0}
    header{padding:24px 32px;border-bottom:1px solid #2a3b4d}
    h1{margin:0 0 6px;font-size:22px}
    .sub{opacity:.8;font-size:14px}
    section{padding:28px 32px;border-bottom:1px solid #22303d}
    h2{margin:0 0 16px;font-size:18px;color:#7eb8da}
    .flow{display:flex;flex-wrap:wrap;gap:10px;align-items:stretch}
    .col{display:flex;flex-direction:column;gap:8px;min-width:180px}
    .box{border-radius:10px;padding:10px 12px;font-size:13px;line-height:1.35;white-space:pre-line;color:#102018}
    .start,.end{background:#e89ab8;border-radius:999px;text-align:center}
    .process{background:#7eb8da}
    .sub{background:#7dbe7d}
    .decision{background:#e89ab8;transform:none;text-align:center}
    .data{background:#e8a05c}
    .ext{background:#d9894a}
    .note{background:#eef3f7}
    .arrow{text-align:center;color:#7eb8da;font-size:18px}
    table{border-collapse:collapse;width:100%;font-size:13px}
    td,th{border:1px solid #2a3b4d;padding:8px 10px;text-align:left}
    th{background:#1a2733}
    """
    parts = [
        "<!DOCTYPE html><html lang='pt-BR'><head><meta charset='utf-8'/>",
        "<title>JudgeTCG — Fluxograma E2E</title>",
        f"<style>{css}</style></head><body>",
        "<header><h1>JudgeTCG — Fluxograma funcional E2E</h1>",
        "<p class='sub'>Modelo Visio (Judge_doc.vsdx) + planilha E2E + arquitetura do repositório. "
        "Arquivo MS Visio: docs/architecture/visio/JudgeTCG_Fluxograma_E2E.vsdx</p></header>",
        "<section><h2>Legenda de formas</h2><table><tr><th>Forma</th><th>Uso</th><th>Público</th></tr>",
        "<tr><td>Início/Fim (rosa)</td><td>Terminador</td><td>Gestores e devs</td></tr>",
        "<tr><td>Processo (azul)</td><td>Atividade da plataforma</td><td>Ambos</td></tr>",
        "<tr><td>Subprocesso (verde)</td><td>Detalhe em outra página</td><td>Ambos</td></tr>",
        "<tr><td>Decisão (losango rosa)</td><td>Regra de negócio / gate</td><td>Gestores</td></tr>",
        "<tr><td>Dados (laranja)</td><td>SQL Postgres / Redis</td><td>Devs</td></tr>",
        "<tr><td>Dados externos (laranja)</td><td>OpenAI, Scryfall, Stripe, BrasilAPI…</td><td>Ambos</td></tr>",
        "</table></section>",
    ]
    for page in pages:
        parts.append(f"<section><h2>Página Visio: {html.escape(page.name)}</h2><div class='flow'>")
        # group roughly by x
        cols: dict[float, list[Node]] = {}
        for n in page.nodes:
            k = round(n.x, 1)
            cols.setdefault(k, []).append(n)
        for k in sorted(cols):
            parts.append("<div class='col'>")
            nodes = sorted(cols[k], key=lambda z: -z.y)
            for i, n in enumerate(nodes):
                txt = html.escape(n.text)
                parts.append(f"<div class='box {n.kind}'>{txt}</div>")
                if i < len(nodes) - 1:
                    parts.append("<div class='arrow'>↓</div>")
            parts.append("</div>")
        parts.append("</div></section>")
    parts.append("</body></html>")
    return "".join(parts)


def write_vsdx(pages: list[Page], dest: Path) -> None:
    if not TEMPLATE.exists():
        raise SystemExit(f"Template Visio não extraído em {TEMPLATE}")
    work = Path("/tmp/judgetcg_vsdx_build")
    if work.exists():
        shutil.rmtree(work)
    shutil.copytree(TEMPLATE, work)
    pages_dir = work / "visio" / "pages"
    rels_dir = pages_dir / "_rels"
    # remove old page xml (keep structure)
    for p in pages_dir.glob("page*.xml"):
        p.unlink()

    page_elems = []
    ct_overrides = []
    rels_pages = []
    titles = []
    for i, page in enumerate(pages):
        pid = i
        fname = f"page{i+1}.xml"
        unique = uid()
        xml, _ = page_contents(page)
        (pages_dir / fname).write_text(xml, encoding="utf-8")
        (rels_dir / f"{fname}.rels").write_text(page_rels(), encoding="utf-8")
        page_elems.append(page_sheet(page, pid, unique))
        rels_pages.append(
            f'<Relationship Id="rId{i+1}" Type="http://schemas.microsoft.com/visio/2010/relationships/page" Target="{fname}"/>'
        )
        ct_overrides.append(
            f'<Override PartName="/visio/pages/{fname}" ContentType="application/vnd.ms-visio.page+xml"/>'
        )
        titles.append(page.name)

    pages_xml = (
        f"<?xml version='1.0' encoding='utf-8' ?>"
        f"<Pages xmlns='{NS}' xmlns:r='{REL_NS}' xml:space='preserve'>"
        + "".join(page_elems)
        + "</Pages>"
    )
    (pages_dir / "pages.xml").write_text(pages_xml, encoding="utf-8")
    (rels_dir / "pages.xml.rels").write_text(
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels_pages)
        + "</Relationships>",
        encoding="utf-8",
    )

    # Content_Types: keep masters + replace pages
    types = work / "[Content_Types].xml"
    base = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Default Extension="png" ContentType="image/png"/>
<Default Extension="emf" ContentType="image/x-emf"/>
<Override PartName="/visio/document.xml" ContentType="application/vnd.ms-visio.drawing.main+xml"/>
<Override PartName="/visio/masters/masters.xml" ContentType="application/vnd.ms-visio.masters+xml"/>
"""
    masters_ct = "".join(
        f'<Override PartName="/visio/masters/master{i}.xml" ContentType="application/vnd.ms-visio.master+xml"/>'
        for i in range(1, 11)
    )
    rest = """<Override PartName="/visio/pages/pages.xml" ContentType="application/vnd.ms-visio.pages+xml"/>
<Override PartName="/visio/windows.xml" ContentType="application/vnd.ms-visio.windows+xml"/>
<Override PartName="/visio/validation.xml" ContentType="application/vnd.ms-visio.validation+xml"/>
<Override PartName="/visio/comments.xml" ContentType="application/vnd.ms-visio.comments+xml"/>
<Override PartName="/visio/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>
</Types>"""
    types.write_text(base + masters_ct + "".join(ct_overrides) + rest, encoding="utf-8")

    n = len(pages)
    titles_xml = "".join(f"<vt:lpstr>{escape(t)}</vt:lpstr>" for t in titles)
    masters_titles = (
        "<vt:lpstr>Início/Término</vt:lpstr><vt:lpstr>Dados</vt:lpstr><vt:lpstr>Decisão</vt:lpstr>"
        "<vt:lpstr>Processo</vt:lpstr><vt:lpstr>Conector dinâmico</vt:lpstr>"
        "<vt:lpstr>Retângulo Arredondado</vt:lpstr><vt:lpstr>Gradiente Central</vt:lpstr>"
        "<vt:lpstr>Subprocesso</vt:lpstr><vt:lpstr>Dados Externos</vt:lpstr>"
        "<vt:lpstr>Referência fora da página</vt:lpstr>"
    )
    app = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
<Template>VERTICALFLOWCHART_M.vstx</Template><Application>Microsoft Visio</Application>
<HeadingPairs><vt:vector size="4" baseType="variant">
<vt:variant><vt:lpstr>Páginas</vt:lpstr></vt:variant><vt:variant><vt:i4>{n}</vt:i4></vt:variant>
<vt:variant><vt:lpstr>Mestres</vt:lpstr></vt:variant><vt:variant><vt:i4>10</vt:i4></vt:variant>
</vt:vector></HeadingPairs>
<TitlesOfParts><vt:vector size="{n+10}" baseType="lpstr">{titles_xml}{masters_titles}</vt:vector></TitlesOfParts>
<AppVersion>16.0000</AppVersion>
</Properties>"""
    (work / "docProps" / "app.xml").write_text(app, encoding="utf-8")
    core = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<dc:title>JudgeTCG Fluxograma E2E</dc:title>
<dc:subject>Arquitetura funcional API, banco e dados externos</dc:subject>
<dc:creator>JudgeTCG Platform Guardian</dc:creator>
<cp:lastModifiedBy>JudgeTCG</cp:lastModifiedBy>
<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-09T00:00:00Z</dcterms:created>
<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-09T00:00:00Z</dcterms:modified>
</cp:coreProperties>"""
    (work / "docProps" / "core.xml").write_text(core, encoding="utf-8")

    doc = (work / "visio" / "document.xml").read_text(encoding="utf-8")
    doc = doc.replace("TopPage='9'", "TopPage='0'")
    (work / "visio" / "document.xml").write_text(doc, encoding="utf-8")
    win = (work / "visio" / "windows.xml").read_text(encoding="utf-8")
    win = win.replace("Page='9'", "Page='0'")
    (work / "visio" / "windows.xml").write_text(win, encoding="utf-8")

    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        dest.unlink()
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        for f in work.rglob("*"):
            if f.is_file():
                z.write(f, f.relative_to(work).as_posix())



if __name__ == "__main__":
    main()
