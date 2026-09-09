#!/usr/bin/env python3
"""Fluxograma Visio do Painel do Vendedor — leitura para investidor.

Fonte de UI: /vendedor/painel (seller-sidebar-nav.ts + VendedorPainelClientLayout).
Stencil: Judge_doc.vsdx em /tmp/visio_vsdx.
Não inventa módulos comingSoon nem liquidez (LPC=0).
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from generate_onboarding_visio import (  # noqa: E402
    OUT_DIR,
    Edge,
    Node,
    Page,
    col,
    drawio_xml,
    html_preview,
    write_vsdx,
)


def build_pages() -> list[Page]:
    pages: list[Page] = []

    pages.append(
        Page(
            "Legenda",
            "Legenda",
            13.2,
            8.6,
            [
                Node("l0", "Início / Fim", "start", 2.2, 7.5, 2.3, 0.5),
                Node("l1", "Capacidade do produto", "process", 5.1, 7.5, 2.6, 0.55),
                Node("l2", "Detalhe em outra página", "sub", 8.1, 7.5, 2.6, 0.55),
                Node("l3", "Decisão / gate", "decision", 11.0, 7.5, 2.2, 0.7),
                Node("l4", "Dado interno (SQL)", "data", 2.2, 6.3, 2.6, 0.65),
                Node("l5", "Parceiro (Stripe, PIX…)", "ext", 5.4, 6.3, 2.6, 0.65),
                Node("l6", "Painel do vendedor\nUI no ar · API Render fora", "note", 9.2, 6.3, 4.0, 0.75),
            ],
            [],
        )
    )

    # Tese para investidor
    n = [
        Node("t0", "Investidor: o que este painel é", "start", 6.6, 10.8, 3.2, 0.55),
        Node(
            "t1",
            "JudgeTCG = marketplace TCG BR\n+ loja CNPJ na oferta (não CPF-seller)",
            "process",
            6.6,
            9.5,
            3.6,
            0.8,
        ),
        Node(
            "t2",
            "Painel /vendedor/painel\nferramenta da loja para publicar oferta\ne cumprir pedido sem a equipe no meio",
            "process",
            6.6,
            8.1,
            3.8,
            0.85,
        ),
        Node(
            "t3",
            "Métrica que importa: LPC\noferta × procura sem intervenção humana",
            "decision",
            6.6,
            6.6,
            3.4,
            0.85,
        ),
        Node("t4", "LPC hoje = 0\nBeta não iniciado", "note", 3.0, 5.2, 2.6, 0.7),
        Node(
            "t5",
            "UI Vercel no ar\nFastAPI/Render fora — KPIs podem ser mock",
            "note",
            10.2,
            5.2,
            3.2,
            0.75,
        ),
        Node("t6", "Ciclo E2E (fluxograma real)", "sub", 6.6, 3.8, 3.0, 0.6, "Ciclo E2E"),
    ]
    e = [
        Edge("t0", "t1"),
        Edge("t1", "t2"),
        Edge("t2", "t3"),
        Edge("t3", "t4", "honestidade"),
        Edge("t3", "t5", "runtime 2026-09"),
        Edge("t3", "t6", "capacidades"),
    ]
    pages.append(Page("Tese investidor", "Tese investidor", 14.0, 12.0, n, e))

    # Gates de entrada (screenshot: banner + empty store)
    n, e = col(
        3.4,
        12.2,
        0.95,
        [
            ("g0", "Lojista abre /vendedor/painel", "start"),
            ("g1", "Login Supabase\n/entrar se anônimo", "ext"),
            ("g2", "GET /api/account/status\nKYC merchant", "ext"),
            ("g3", "Status da conta leu?", "decision"),
        ],
    )
    n += [
        Node(
            "g4",
            "Banner vermelho\nStatus indisponível\nO painel permanece disponível",
            "process",
            7.2,
            9.35,
            2.8,
            0.9,
        ),
        Node("g5", "KYC bloqueia publicar?\n/loja/suspensa", "decision", 3.4, 7.4, 2.6, 0.8),
        Node("g6", "GET /api/stores/mine\nHá loja cadastrada?", "decision", 3.4, 6.1, 2.6, 0.8),
        Node("g7", "Empty state\nCadastrar loja", "sub", 7.2, 6.1, 2.5, 0.7, "Cadastrar loja"),
        Node("g8", "Dashboard com KPIs", "sub", 3.4, 4.8, 2.5, 0.6, "Mapa módulos"),
        Node(
            "g9",
            "Fail-open 12s: painel abre\nmesmo se auth/KYC atrasar",
            "note",
            7.2,
            4.8,
            2.8,
            0.75,
        ),
        Node("g10", "Fim entrada", "end", 3.4, 3.5, 2.2, 0.5),
    ]
    e += [
        Edge("g3", "g4", "Não (503 API)", "#e51400"),
        Edge("g3", "g5", "Sim", "#75d175"),
        Edge("g4", "g6"),
        Edge("g5", "g6", "pode publicar"),
        Edge("g6", "g7", "Não", "#e51400"),
        Edge("g6", "g8", "Sim", "#75d175"),
        Edge("g8", "g10"),
        Edge("g7", "g10"),
    ]
    pages.append(Page("Entrada e gates", "Entrada e gates", 12.0, 13.4, n, e))

    # Mapa de módulos (sidebar do print)
    n = [
        Node("m0", "Painel do vendedor", "start", 7.0, 12.4, 2.6, 0.5),
        Node("m1", "Dashboard · Operação · Inbox", "process", 2.6, 11.0, 2.8, 0.7),
        Node("m2", "Insights (plano analytics)", "process", 6.0, 11.0, 2.6, 0.7),
        Node("m3", "Anúncios + Novo anúncio", "sub", 9.4, 11.0, 2.6, 0.7, "Publicar oferta"),
        Node("m4", "Catálogo\ncartas · produtos · expansões · jogos", "process", 2.6, 9.4, 2.8, 0.8),
        Node("m5", "Estoque + import CSV", "process", 6.0, 9.4, 2.6, 0.7),
        Node("m6", "Pedidos (pipeline)\nbadge = pendentes de pagamento", "sub", 9.4, 9.4, 2.8, 0.8, "Pedidos"),
        Node("m7", "Clientes (lista / CRM plano)", "process", 2.6, 7.8, 2.8, 0.7),
        Node("m8", "Atendimento tickets\nbadge = tickets abertos", "process", 6.0, 7.8, 2.6, 0.75),
        Node("m9", "Financeiro\nreceitas · repasses · Stripe · PIX", "sub", 9.4, 7.8, 2.8, 0.8, "Financeiro"),
        Node("m10", "Marketing: cupons · buylist · ingressos", "process", 2.6, 6.2, 2.9, 0.75),
        Node("m11", "Reputação + Estatísticas", "process", 6.0, 6.2, 2.6, 0.7),
        Node("m12", "Equipe + Configurações + PDV", "process", 9.4, 6.2, 2.8, 0.75),
        Node("m13", "Ver loja pública  ·  Voltar ao site", "process", 6.0, 4.7, 3.2, 0.65),
        Node(
            "m14",
            "Ocultos (comingSoon): chat, CRM avançado,\nreclamações, promoções, movimentações",
            "note",
            6.0,
            3.4,
            4.0,
            0.75,
        ),
        Node("m15", "Ciclo E2E (fluxograma real)", "sub", 6.0, 2.1, 3.2, 0.6, "Ciclo E2E"),
        Node("m16", "Dashboard detalhe", "sub", 2.6, 2.1, 2.4, 0.55, "Dashboard"),
        Node("m17", "Estoque detalhe", "sub", 9.8, 2.1, 2.4, 0.55, "Estoque"),
    ]
    e = [
        Edge("m0", "m1"),
        Edge("m0", "m2"),
        Edge("m0", "m3"),
        Edge("m0", "m4"),
        Edge("m0", "m5"),
        Edge("m0", "m6"),
        Edge("m4", "m3", "publica"),
        Edge("m5", "m6", "atende pedido"),
        Edge("m6", "m9", "repassa $"),
        Edge("m0", "m7"),
        Edge("m0", "m8"),
        Edge("m0", "m10"),
        Edge("m0", "m11"),
        Edge("m0", "m12"),
        Edge("m0", "m13"),
        Edge("m13", "m15"),
        Edge("m1", "m16"),
        Edge("m5", "m17"),
    ]
    pages.append(Page("Mapa módulos", "Mapa módulos", 13.4, 13.8, n, e))

    # ---- Ciclo E2E (fluxograma real para investidor) ----
    n = [
        Node("e00", "Início: lojista quer vender no JudgeTCG", "start", 3.4, 22.6, 3.0, 0.55),
        Node("e01", "GET /entrar · Supabase Auth", "ext", 3.4, 21.5, 2.7, 0.6),
        Node("e02", "Sessão JWT?", "decision", 3.4, 20.4, 2.4, 0.7),
        Node("e03", "Redirect /entrar?next=painel", "end", 7.0, 20.4, 2.6, 0.6),
        Node("e04", "GET /api/account/status\nKYC merchant + player", "ext", 3.4, 19.2, 2.8, 0.7),
        Node("e05", "Status leu?", "decision", 3.4, 18.05, 2.3, 0.65),
        Node("e06", "Banner: status indisponível\npainel fail-open", "process", 7.0, 18.05, 2.7, 0.75),
        Node("e07", "can_publish / KYC ok?", "decision", 3.4, 16.85, 2.5, 0.7),
        Node("e08", "GET /loja/suspensa\nnão publica oferta", "end", 7.0, 16.85, 2.6, 0.7),
        Node("e09", "GET /api/stores/mine", "data", 3.4, 15.7, 2.6, 0.6),
        Node("e10", "Há loja?", "decision", 3.4, 14.55, 2.3, 0.65),
        Node("e11", "POST /api/stores\n/stores/create + CNPJ local", "sub", 7.0, 14.55, 2.7, 0.75, "Cadastrar loja"),
        Node("e12", "POST Connect onboard\nStripe conta da loja", "ext", 10.4, 14.55, 2.6, 0.75),
        Node("e13", "Dashboard GET /api/seller/dashboard/overview", "process", 3.4, 13.3, 3.0, 0.7),
        Node("e14", "Overview 200?", "decision", 3.4, 12.1, 2.3, 0.65),
        Node("e15", "Mock KPIs\npending_payment=12 tickets=3", "note", 7.0, 12.1, 2.7, 0.7),
        Node("e16", "Catálogo GET cartas/produtos/jogos", "sub", 3.4, 10.9, 2.8, 0.7, "Catálogo"),
        Node("e17", "Wizard /listagens/nova\nPOST /api/seller/listings", "sub", 3.4, 9.7, 2.8, 0.75, "Publicar oferta"),
        Node("e18", "Estoque adjust/CSV\nPOST /api/seller/inventory/*", "sub", 7.0, 9.7, 2.8, 0.75, "Estoque"),
        Node("e19", "Buyer vê PDP / listing", "process", 3.4, 8.5, 2.6, 0.6),
        Node("e20", "Checkout Stripe ou PIX", "ext", 3.4, 7.4, 2.6, 0.6),
        Node("e21", "Pedido SQL marketplace", "data", 3.4, 6.3, 2.6, 0.6),
        Node("e22", "Loja: tabs Pedidos\nGET /api/seller/orders", "sub", 3.4, 5.15, 2.7, 0.7, "Pedidos"),
        Node("e23", "Fulfillment bulk commands\nseparar / enviar", "data", 7.0, 5.15, 2.7, 0.75),
        Node("e24", "Ticket GET /api/seller/tickets", "process", 10.4, 5.15, 2.6, 0.7),
        Node("e25", "Repasse Stripe Connect / PIX", "sub", 3.4, 3.95, 2.8, 0.7, "Financeiro"),
        Node("e26", "LPC possível?\nsem equipe JudgeTCG", "decision", 3.4, 2.75, 2.6, 0.75),
        Node("e27", "LPC = 0 hoje\nBeta not started", "end", 7.0, 2.75, 2.6, 0.65),
        Node("e28", "Ciclo de produto desenhado", "end", 3.4, 1.6, 2.6, 0.55),
    ]
    e = [
        Edge("e00", "e01"),
        Edge("e01", "e02"),
        Edge("e02", "e03", "Não", "#e51400"),
        Edge("e02", "e04", "Sim", "#75d175"),
        Edge("e04", "e05"),
        Edge("e05", "e06", "503", "#e51400"),
        Edge("e05", "e07", "Sim", "#75d175"),
        Edge("e06", "e09"),
        Edge("e07", "e08", "Não", "#e51400"),
        Edge("e07", "e09", "Sim", "#75d175"),
        Edge("e09", "e10"),
        Edge("e10", "e11", "Não", "#e51400"),
        Edge("e10", "e13", "Sim", "#75d175"),
        Edge("e11", "e12"),
        Edge("e11", "e13"),
        Edge("e13", "e14"),
        Edge("e14", "e15", "Não", "#e51400"),
        Edge("e14", "e16", "Sim", "#75d175"),
        Edge("e15", "e16"),
        Edge("e16", "e17"),
        Edge("e17", "e18"),
        Edge("e17", "e19"),
        Edge("e19", "e20"),
        Edge("e20", "e21", "SQL"),
        Edge("e21", "e22"),
        Edge("e22", "e23"),
        Edge("e22", "e24", "disputa"),
        Edge("e23", "e25"),
        Edge("e25", "e26"),
        Edge("e26", "e27", "ainda não", "#e51400"),
        Edge("e26", "e28", "quando houver evidência", "#75d175"),
    ]
    pages.append(Page("Ciclo E2E", "Ciclo E2E", 14.0, 24.0, n, e))

    # ---- Dashboard ----
    n, e = col(
        3.3,
        14.2,
        0.85,
        [
            ("db0", "GET /vendedor/painel", "start"),
            ("db1", "useSellerStore · stores/mine", "data"),
            ("db2", "Há loja?", "decision"),
        ],
    )
    n += [
        Node("db3", "Empty: Cadastrar loja\n+ MerchantKycCard", "sub", 7.0, 12.5, 2.6, 0.75, "Cadastrar loja"),
        Node("db4", "GET /api/seller/dashboard/overview", "ext", 3.3, 10.8, 2.8, 0.7),
        Node("db5", "200?", "decision", 3.3, 9.7, 2.2, 0.6),
        Node("db6", "dashboardOverviewMock()", "note", 7.0, 9.7, 2.5, 0.6),
        Node("db7", "KPI strip · SLA · low stock\npedidos recentes · tickets", "process", 3.3, 8.5, 2.8, 0.75),
        Node("db8", "SellerOpsDashboardV2\ncommand center", "process", 3.3, 7.3, 2.7, 0.7),
        Node("db9", "GET account status → KYC card", "ext", 3.3, 6.15, 2.7, 0.65),
        Node("db10", "Quick actions: novo anúncio\nver loja pública", "process", 3.3, 5.0, 2.7, 0.7),
        Node("db11", "Fim dashboard", "end", 3.3, 3.85, 2.2, 0.5),
    ]
    e += [
        Edge("db2", "db3", "Não", "#e51400"),
        Edge("db2", "db4", "Sim", "#75d175"),
        Edge("db4", "db5"),
        Edge("db5", "db6", "Não", "#e51400"),
        Edge("db5", "db7", "Sim", "#75d175"),
        Edge("db6", "db7"),
        Edge("db7", "db8"),
        Edge("db8", "db9"),
        Edge("db9", "db10"),
        Edge("db10", "db11"),
        Edge("db3", "db11"),
    ]
    pages.append(Page("Dashboard", "Dashboard", 11.4, 15.4, n, e))

    # ---- Operação + Inbox + Insights ----
    n = [
        Node("op0", "Operação diária", "start", 3.2, 10.8, 2.4, 0.5),
        Node("op1", "GET /vendedor/painel/operacao\nfilas overview", "process", 3.2, 9.6, 2.6, 0.7),
        Node("op2", "Inbox\nnotificações + ações", "process", 6.4, 9.6, 2.5, 0.7),
        Node("op3", "Insights\nGET /api/seller/ai/insights", "ext", 9.6, 9.6, 2.6, 0.75),
        Node("op4", "Plano analytics?", "decision", 9.6, 8.3, 2.3, 0.7),
        Node("op5", "Lock → /vendedor/painel/planos", "process", 9.6, 7.1, 2.6, 0.65),
        Node("op6", "Chargebacks / tickets na fila", "process", 3.2, 8.3, 2.6, 0.7),
        Node("op7", "Abrir Pedidos ou Atendimento", "sub", 3.2, 7.1, 2.6, 0.65, "Pedidos"),
        Node("op8", "Fim ops", "end", 3.2, 5.9, 2.0, 0.45),
    ]
    e = [
        Edge("op0", "op1"),
        Edge("op0", "op2"),
        Edge("op0", "op3"),
        Edge("op3", "op4"),
        Edge("op4", "op5", "Não", "#e51400"),
        Edge("op1", "op6"),
        Edge("op6", "op7"),
        Edge("op7", "op8"),
    ]
    pages.append(Page("Operação Inbox", "Operação Inbox", 13.0, 12.0, n, e))

    # ---- Catálogo ----
    n = [
        Node("k0", "Loja precisa de SKU", "start", 3.2, 11.4, 2.4, 0.5),
        Node("k1", "Tipo de catálogo?", "decision", 3.2, 10.2, 2.4, 0.7),
        Node("k2", "Cartas\n/catalogo/cartas", "process", 6.4, 10.2, 2.4, 0.65),
        Node("k3", "Produtos\n/catalogo/produtos", "process", 9.2, 10.2, 2.4, 0.65),
        Node("k4", "Expansões / Jogos", "process", 12.0, 10.2, 2.3, 0.65),
        Node("k5", "GET catálogo FastAPI\n(search / games)", "ext", 3.2, 8.7, 2.6, 0.7),
        Node("k6", "API ok?", "decision", 3.2, 7.5, 2.2, 0.65),
        Node("k7", "Mock de cartas/jogos\n(degradação)", "note", 6.4, 7.5, 2.4, 0.65),
        Node("k8", "Seleciona carta/produto", "process", 3.2, 6.3, 2.5, 0.6),
        Node("k9", "Publicar listing", "sub", 3.2, 5.15, 2.4, 0.6, "Publicar oferta"),
        Node("k10", "Fim catálogo", "end", 3.2, 4.05, 2.1, 0.45),
    ]
    e = [
        Edge("k0", "k1"),
        Edge("k1", "k2", "single"),
        Edge("k1", "k3", "sealed/acessório"),
        Edge("k1", "k4", "set/TCG"),
        Edge("k2", "k5"),
        Edge("k3", "k5"),
        Edge("k4", "k5"),
        Edge("k5", "k6"),
        Edge("k6", "k7", "Não", "#e51400"),
        Edge("k6", "k8", "Sim", "#75d175"),
        Edge("k7", "k8"),
        Edge("k8", "k9"),
        Edge("k9", "k10"),
    ]
    pages.append(Page("Catálogo", "Catálogo", 15.0, 12.6, n, e))

    # ---- Estoque ----
    n, e = col(
        3.3,
        12.4,
        0.88,
        [
            ("st0", "GET /vendedor/painel/estoque", "start"),
            ("st1", "Há loja?", "decision"),
        ],
    )
    n += [
        Node("st2", "Empty cadastrar loja", "end", 6.8, 11.52, 2.4, 0.55),
        Node("st3", "GET inventory search", "ext", 3.3, 10.4, 2.5, 0.6),
        Node("st4", "200?", "decision", 3.3, 9.3, 2.2, 0.6),
        Node("st5", "HTTP 503 (sem mock)", "end", 6.8, 9.3, 2.4, 0.55),
        Node("st6", "Ajuste qty\nPOST /api/seller/inventory/adjust", "data", 3.3, 8.1, 2.8, 0.7),
        Node("st7", "Bulk POST /inventory/bulk", "data", 6.8, 8.1, 2.5, 0.65),
        Node("st8", "CSV POST /inventory/import-csv", "data", 3.3, 6.9, 2.8, 0.7),
        Node("st9", "GET /inventory/export", "ext", 6.8, 6.9, 2.4, 0.6),
        Node("st10", "Low stock → widget dashboard", "process", 3.3, 5.7, 2.6, 0.65),
        Node("st11", "comingSoon: movimentações\nfora do menu", "note", 6.8, 5.7, 2.5, 0.7),
        Node("st12", "Fim estoque", "end", 3.3, 4.5, 2.1, 0.45),
    ]
    e += [
        Edge("st1", "st2", "Não", "#e51400"),
        Edge("st1", "st3", "Sim", "#75d175"),
        Edge("st3", "st4"),
        Edge("st4", "st5", "Não", "#e51400"),
        Edge("st4", "st6", "Sim", "#75d175"),
        Edge("st6", "st7"),
        Edge("st6", "st8"),
        Edge("st8", "st9"),
        Edge("st6", "st10"),
        Edge("st10", "st12"),
    ]
    pages.append(Page("Estoque", "Estoque", 11.2, 13.6, n, e))

    # ---- Pedidos (tabs reais do menu) ----
    n = [
        Node("pd0", "GET /vendedor/painel/pedidos", "start", 3.2, 13.6, 2.6, 0.5),
        Node("pd1", "GET /api/seller/orders", "ext", 3.2, 12.5, 2.5, 0.6),
        Node("pd2", "200?", "decision", 3.2, 11.4, 2.2, 0.6),
        Node("pd3", "HTTP 503 (sem mock)", "end", 6.8, 11.4, 2.4, 0.55),
        Node("pd4", "Tab?", "decision", 3.2, 10.2, 2.2, 0.65),
        Node("pd5", "pending_payment", "process", 6.2, 10.2, 2.2, 0.55),
        Node("pd6", "paid / to_separate", "process", 8.7, 10.2, 2.3, 0.55),
        Node("pd7", "shipped / delivered", "process", 11.3, 10.2, 2.3, 0.55),
        Node("pd8", "cancelled / refunded", "process", 6.2, 9.0, 2.3, 0.55),
        Node("pd9", "GET /api/seller/orders/{id}\nOrderDetailDrawer", "data", 3.2, 9.0, 2.6, 0.7),
        Node("pd10", "POST bulk/fulfillment/commands", "data", 3.2, 7.7, 2.7, 0.7),
        Node("pd11", "Melhor Envio cotação (se frete)", "ext", 6.8, 7.7, 2.6, 0.7),
        Node("pd12", "Atualiza status no Postgres", "data", 3.2, 6.5, 2.6, 0.6),
        Node("pd13", "Badge sidebar =\noverview.pending_payment", "note", 6.8, 6.5, 2.6, 0.7),
        Node("pd14", "Segue para Financeiro", "sub", 3.2, 5.3, 2.5, 0.6, "Financeiro"),
        Node("pd15", "Fim pedidos", "end", 3.2, 4.15, 2.1, 0.45),
    ]
    e = [
        Edge("pd0", "pd1"),
        Edge("pd1", "pd2"),
        Edge("pd2", "pd3", "Não", "#e51400"),
        Edge("pd2", "pd4", "Sim", "#75d175"),
        Edge("pd4", "pd5", "pagar"),
        Edge("pd4", "pd6", "separar"),
        Edge("pd4", "pd7", "trânsito"),
        Edge("pd4", "pd8", "exceção"),
        Edge("pd5", "pd9"),
        Edge("pd6", "pd9"),
        Edge("pd7", "pd9"),
        Edge("pd9", "pd10"),
        Edge("pd10", "pd11"),
        Edge("pd10", "pd12"),
        Edge("pd12", "pd14"),
        Edge("pd14", "pd15"),
    ]
    pages.append(Page("Pedidos", "Pedidos", 14.4, 14.8, n, e))

    # ---- Financeiro detalhado ----
    n = [
        Node("fn0", "GET /financeiro/receitas", "start", 3.2, 12.2, 2.6, 0.5),
        Node("fn1", "De onde veio o $?", "decision", 3.2, 11.0, 2.4, 0.7),
        Node("fn2", "Stripe Checkout/Connect", "ext", 6.4, 11.0, 2.5, 0.65),
        Node("fn3", "PIX chave da loja", "ext", 9.4, 11.0, 2.4, 0.6),
        Node("fn4", "GET /financeiro/repasses", "process", 3.2, 9.6, 2.6, 0.6),
        Node("fn5", "GET /financeiro/stripe", "ext", 6.4, 9.6, 2.5, 0.6),
        Node("fn6", "GET /financeiro/pix", "ext", 9.4, 9.6, 2.4, 0.6),
        Node("fn7", "Reconciliação", "process", 3.2, 8.3, 2.4, 0.55),
        Node("fn8", "Chargebacks", "process", 6.4, 8.3, 2.4, 0.55),
        Node("fn9", "Auditoria trail", "process", 9.4, 8.3, 2.4, 0.55),
        Node("fn10", "API finance 200?", "decision", 3.2, 7.05, 2.4, 0.65),
        Node("fn11", "Mock financeiro\n(não é receita real)", "note", 6.6, 7.05, 2.6, 0.7),
        Node("fn12", "Payout na conta da loja", "end", 3.2, 5.8, 2.5, 0.55),
    ]
    e = [
        Edge("fn0", "fn1"),
        Edge("fn1", "fn2", "cartão"),
        Edge("fn1", "fn3", "PIX"),
        Edge("fn2", "fn4"),
        Edge("fn3", "fn6"),
        Edge("fn4", "fn5"),
        Edge("fn4", "fn7"),
        Edge("fn5", "fn8"),
        Edge("fn7", "fn9"),
        Edge("fn7", "fn10"),
        Edge("fn10", "fn11", "Não", "#e51400"),
        Edge("fn10", "fn12", "Sim", "#75d175"),
        Edge("fn11", "fn12"),
    ]
    pages.append(Page("Financeiro", "Financeiro", 13.0, 13.4, n, e))

    # ---- Marketing / clientes / atendimento / reputação ----
    n = [
        Node("mk0", "Crescimento da loja", "start", 6.6, 11.6, 2.6, 0.5),
        Node("mk1", "GET /cupons\nCRUD cupom da loja", "process", 2.6, 10.2, 2.5, 0.7),
        Node("mk2", "GET /buylist\nplano buylist", "process", 6.6, 10.2, 2.5, 0.7),
        Node("mk3", "GET /ingressos\nplano tournaments", "process", 10.6, 10.2, 2.6, 0.7),
        Node("mk4", "Plano libera?", "decision", 6.6, 8.8, 2.4, 0.7),
        Node("mk5", "Upsell /planos", "process", 10.6, 8.8, 2.4, 0.6),
        Node("mk6", "Clientes /lista\nGET /api/seller/customers", "ext", 2.6, 7.5, 2.6, 0.75),
        Node("mk7", "Tickets\nGET /api/seller/tickets", "ext", 6.6, 7.5, 2.5, 0.7),
        Node("mk8", "Reputação\nGET /api/seller/reputation", "ext", 10.6, 7.5, 2.6, 0.75),
        Node("mk9", "comingSoon fora do menu:\nchat · CRM · reclamações · promoções", "note", 6.6, 6.1, 3.6, 0.7),
        Node("mk10", "Fim marketing", "end", 6.6, 4.9, 2.2, 0.45),
    ]
    e = [
        Edge("mk0", "mk1"),
        Edge("mk0", "mk2"),
        Edge("mk0", "mk3"),
        Edge("mk2", "mk4"),
        Edge("mk3", "mk4"),
        Edge("mk4", "mk5", "Não", "#e51400"),
        Edge("mk4", "mk6", "Sim", "#75d175"),
        Edge("mk1", "mk6"),
        Edge("mk6", "mk7"),
        Edge("mk7", "mk8"),
        Edge("mk8", "mk10"),
    ]
    pages.append(Page("Marketing CRM", "Marketing CRM", 14.2, 12.8, n, e))

    # ---- Equipe config PDV ----
    n = [
        Node("eq0", "Operar a loja como time", "start", 3.4, 11.4, 2.6, 0.5),
        Node("eq1", "Equipe usuários\nGET team/users · convites", "ext", 3.4, 10.2, 2.7, 0.7),
        Node("eq2", "Permissões (matriz role)", "process", 6.6, 10.2, 2.5, 0.65),
        Node("eq3", "Logs GET team/logs", "data", 9.6, 10.2, 2.4, 0.6),
        Node("eq4", "Configurações loja\nPUT /api/stores/id/{id}", "data", 3.4, 8.8, 2.7, 0.7),
        Node("eq5", "/configuracoes/pagamentos", "sub", 6.6, 8.8, 2.6, 0.65, "Cadastrar loja"),
        Node("eq6", "/configuracoes/frete\nMelhor Envio", "ext", 9.6, 8.8, 2.5, 0.7),
        Node("eq7", "Notificações settings", "process", 3.4, 7.5, 2.6, 0.6),
        Node("eq8", "PDV plano Pro+\n/pdv + produtos locais", "decision", 6.6, 7.5, 2.6, 0.75),
        Node("eq9", "Lock /planos", "end", 9.8, 7.5, 2.2, 0.5),
        Node("eq10", "Balcão / PIX local", "process", 6.6, 6.2, 2.5, 0.6),
        Node("eq11", "Fim backoffice", "end", 3.4, 6.2, 2.2, 0.5),
    ]
    e = [
        Edge("eq0", "eq1"),
        Edge("eq1", "eq2"),
        Edge("eq2", "eq3"),
        Edge("eq0", "eq4"),
        Edge("eq4", "eq5"),
        Edge("eq4", "eq6"),
        Edge("eq4", "eq7"),
        Edge("eq0", "eq8"),
        Edge("eq8", "eq9", "sem plano", "#e51400"),
        Edge("eq8", "eq10", "Pro+", "#75d175"),
        Edge("eq7", "eq11"),
        Edge("eq10", "eq11"),
    ]
    pages.append(Page("Equipe PDV", "Equipe PDV", 13.2, 12.6, n, e))

    # Cadastrar loja (honest dual path)
    n, e = col(
        3.3,
        11.4,
        0.95,
        [
            ("c0", "CTA Cadastrar loja", "start"),
            ("c1", "Empty state do dashboard\n(sem GET /api/stores/mine)", "process"),
            ("c2", "Para onde o botão leva?", "decision"),
        ],
    )
    n += [
        Node(
            "c3",
            "/vendedor/painel/onboarding\nStripe Connect\nNÃO cria a loja",
            "ext",
            7.0,
            8.55,
            2.8,
            0.9,
        ),
        Node(
            "c4",
            "/stores/create\nPOST /api/stores\ncria a loja de verdade",
            "data",
            3.3,
            6.6,
            2.7,
            0.85,
        ),
        Node("c5", "Configurações / pagamentos", "process", 3.3, 5.4, 2.6, 0.65),
        Node(
            "c6",
            "Investidor: onboarding de oferta\n= entidade loja + CNPJ + Connect",
            "note",
            7.0,
            5.4,
            2.9,
            0.8,
        ),
        Node("c7", "Loja pronta para anunciar", "end", 3.3, 4.2, 2.4, 0.55),
    ]
    e += [
        Edge("c2", "c3", "atalho empty state"),
        Edge("c2", "c4", "cadastro real"),
        Edge("c4", "c5"),
        Edge("c3", "c5", "se já houver loja"),
        Edge("c5", "c7"),
    ]
    pages.append(Page("Cadastrar loja", "Cadastrar loja", 11.6, 12.6, n, e))

    # Publicar oferta
    n, e = col(
        3.3,
        12.0,
        0.92,
        [
            ("p0", "+ Novo anúncio", "start"),
            ("p1", "Header → catálogo/cartas\n(?action=new ainda não abre wizard)", "process"),
            ("p2", "Caminho completo:\n/listagens/nova", "process"),
            ("p3", "GET catálogo cartas (search)", "ext"),
            ("p4", "ListingPublishWizard\npreço, qty, condição, idioma", "process"),
            ("p5", "POST /api/seller/listings", "data"),
            ("p6", "Limite de plano / KYC / loja?", "decision"),
        ],
    )
    n += [
        Node("p7", "Anúncio no ar\nGET /listagens + PDP buyer", "end", 3.3, 4.9, 2.6, 0.7),
        Node("p8", "Bloqueio /loja/suspensa\nou upsell /planos", "end", 7.2, 6.35, 2.6, 0.75),
        Node("p9", "PATCH/DELETE /api/seller/listings/{id}", "data", 7.2, 5.1, 2.7, 0.7),
        Node(
            "p10",
            "Sem FastAPI: listagens 503\nCatálogo pode mock; wizard não grava",
            "note",
            7.2,
            3.8,
            2.8,
            0.8,
        ),
    ]
    e += [
        Edge("p6", "p7", "ok", "#75d175"),
        Edge("p6", "p8", "não", "#e51400"),
        Edge("p7", "p9", "editar depois"),
    ]
    pages.append(Page("Publicar oferta", "Publicar oferta", 11.8, 13.6, n, e))

    # Pedido ao dinheiro
    n = [
        Node("d0", "Buyer paga no marketplace", "start", 3.2, 11.6, 2.5, 0.55),
        Node("d1", "Pedido entra no pipeline", "process", 3.2, 10.4, 2.5, 0.6),
        Node("d2", "Aguardando pagamento → Pago", "process", 3.2, 9.2, 2.6, 0.6),
        Node("d3", "Separação → Enviado → Entregue", "process", 3.2, 8.0, 2.6, 0.65),
        Node("d4", "Tickets se houver disputa", "process", 6.6, 8.0, 2.4, 0.65),
        Node("d5", "Stripe / PIX / Connect", "ext", 3.2, 6.6, 2.5, 0.65),
        Node("d6", "Repasses ao lojista", "ext", 6.6, 6.6, 2.4, 0.6),
        Node("d7", "Chargeback / reconciliação", "process", 9.8, 6.6, 2.5, 0.7),
        Node(
            "d8",
            "LPC: este ciclo sem a equipe JudgeTCG",
            "note",
            6.6,
            5.2,
            3.0,
            0.7,
        ),
        Node("d9", "Dinheiro na loja", "end", 3.2, 5.2, 2.2, 0.5),
        Node(
            "d10",
            "Badges 12 / 3 no print:\nmock se overview API falhar",
            "note",
            9.8,
            8.0,
            2.5,
            0.75,
        ),
    ]
    e = [
        Edge("d0", "d1"),
        Edge("d1", "d2"),
        Edge("d2", "d3"),
        Edge("d3", "d4", "suporte"),
        Edge("d3", "d5"),
        Edge("d5", "d6"),
        Edge("d6", "d7"),
        Edge("d6", "d9"),
        Edge("d8", "d9"),
    ]
    pages.append(Page("Pedido ao $", "Pedido ao $", 13.4, 12.8, n, e))

    # Due diligence / o que está no ar
    n = [
        Node("v0", "Due diligence — o que está no ar", "start", 6.8, 10.6, 3.4, 0.55),
        Node(
            "v1",
            "Pronto na UI (Vercel)\nshell, nav, empty states, wizard, PDV, cupons…",
            "process",
            3.2,
            8.8,
            3.2,
            0.85,
        ),
        Node(
            "v2",
            "Precisa FastAPI + Postgres\nlistagens, pedidos, estoque, stats",
            "ext",
            7.0,
            8.8,
            3.0,
            0.85,
        ),
        Node(
            "v3",
            "Degrada com mock\ndashboard, financeiro, tickets, reputação",
            "note",
            10.8,
            8.8,
            3.0,
            0.85,
        ),
        Node(
            "v4",
            "Não está no menu (comingSoon)\nchat · CRM · promoções · movimentações",
            "process",
            3.2,
            7.0,
            3.2,
            0.8,
        ),
        Node(
            "v5",
            "Não vender como liquidez:\nseed ≠ loja; mock ≠ pedido real",
            "decision",
            7.0,
            7.0,
            3.0,
            0.8,
        ),
        Node(
            "v6",
            "Decisão de orçamento:\n1 runtime FastAPI + API_PROXY_TARGET",
            "note",
            10.8,
            7.0,
            3.0,
            0.8,
        ),
        Node("v7", "Tese intacta: software de loja TCG\npara fechar LPC com CNPJ", "end", 7.0, 5.4, 3.4, 0.7),
    ]
    e = [
        Edge("v0", "v1"),
        Edge("v0", "v2"),
        Edge("v0", "v3"),
        Edge("v1", "v5"),
        Edge("v2", "v5"),
        Edge("v3", "v5"),
        Edge("v4", "v5"),
        Edge("v6", "v7"),
        Edge("v5", "v7", "seguir com evidência"),
    ]
    pages.append(Page("Due diligence", "Due diligence", 14.6, 11.8, n, e))
    return pages


def write_xlsx(dest: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    fill = PatternFill("solid", fgColor="1B4F72")
    font = Font(color="FFFFFF", bold=True)

    ws = wb.active
    ws.title = "Modulos_Investidor"
    headers = [
        "Módulo (menu)",
        "Rota",
        "O que o investidor deve entender",
        "Status no código",
        "Sem FastAPI/Render",
    ]
    ws.append(headers)
    rows = [
        ("Dashboard", "/vendedor/painel", "Centro da operação da loja", "UI live", "KPIs mock; empty se stores/mine 503"),
        ("Operação", "/vendedor/painel/operacao", "Filas: pedidos, tickets, chargebacks", "UI live", "mock se overview cair"),
        ("Inbox", "/vendedor/painel/inbox", "Alertas operacionais unificados", "UI live", "mock"),
        ("Insights", "/vendedor/painel/insights", "Sugestões (plano analytics)", "UI + lock plano", "mock / lock"),
        ("Anúncios", "/vendedor/painel/listagens", "CRUD da oferta no marketplace", "UI live", "503 sem mock"),
        ("Catálogo", "/vendedor/painel/catalogo/*", "Master de cartas/produtos/jogos", "UI live", "catálogo pode mock"),
        ("Estoque", "/vendedor/painel/estoque", "ERP: ajuste, CSV, export", "UI live", "503; empty sem loja"),
        ("Pedidos", "/vendedor/painel/pedidos", "Pipeline pago→envio (núcleo LPC)", "UI live", "503; badge 12 é mock"),
        ("Clientes", "/vendedor/painel/clientes/lista", "CRM básico (plano crm)", "UI live", "mock lista"),
        ("Atendimento", "/vendedor/painel/atendimento/tickets", "Tickets da loja", "UI live", "badge 3 é mock"),
        ("Financeiro", "/vendedor/painel/financeiro/*", "Receita, repasse, Stripe, PIX", "UI live", "mock"),
        ("Cupons", "/vendedor/painel/cupons", "Promoção da loja", "UI live", "precisa loja+API"),
        ("Buylist", "/vendedor/painel/buylist", "Loja compra da comunidade", "UI + plano", "precisa loja"),
        ("Ingressos", "/vendedor/painel/ingressos", "Eventos/torneios da loja", "UI + plano", "precisa loja"),
        ("Reputação", "/vendedor/painel/reputacao", "Trust / SLA / reviews", "UI live", "mock"),
        ("Estatísticas", "/vendedor/painel/estatisticas", "Vendas (plano analytics)", "UI live", "503"),
        ("Equipe", "/vendedor/painel/equipe/*", "Papéis e logs da loja", "UI live", "mock usuários"),
        ("Configurações", "/vendedor/painel/configuracoes", "Loja, frete, pagamentos", "UI live", "503 settings"),
        ("PDV", "/vendedor/painel/pdv", "Balcão físico (plano Pro+)", "UI live", "precisa loja+API"),
        ("Cadastrar loja (empty)", "/vendedor/painel/onboarding", "Atalho = Stripe Connect, não cria loja", "UI live", "Connect precisa API"),
        ("Cadastrar loja (real)", "/stores/create", "POST /api/stores cria entidade", "UI live", "precisa API"),
        ("+ Novo anúncio", "/vendedor/painel/listagens/nova", "Wizard de listing", "UI live", "precisa API"),
        ("Chat / CRM / Promoções", "—", "comingSoon: fora do menu", "stub oculto", "não vender"),
    ]
    for r in rows:
        ws.append(list(r))
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True)
    for i, w in enumerate([28, 42, 48, 22, 38], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = ws.max_row
    tab = Table(displayName="ModulosInvestidor", ref=f"A1:E{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"

    how = wb.create_sheet("Como_importar_Visio")
    how.append(["Passo", "Ação"])
    for row in [
        ("1", "Abrir JudgeTCG_Painel_Vendedor_Investidor.vsdx no Microsoft Visio."),
        ("2", "Abas: Tese investidor → Entrada e gates → Mapa módulos → fluxos → Due diligence."),
        ("3", "Aba Modulos_Investidor deste Excel: copiar linhas para formas no stencil Fluxograma Básico."),
        ("4", "Não tratar badges 12/3 nem empty state como prova de tração (são mock/API fora)."),
        ("6", "Abas novas do .vsdx: Ciclo E2E (fluxograma mestre), Dashboard, Operação Inbox, Catálogo, Estoque, Pedidos, Financeiro, Marketing CRM, Equipe PDV."),
        ("7", "Abas F01_Ciclo_E2E neste Excel: Visualizador de Dados (Passo ID → próximo)."),
    ]:
        how.append(list(row))
    for cell in how[1]:
        cell.fill = fill
        cell.font = font
    how.column_dimensions["A"].width = 8
    how.column_dimensions["B"].width = 110

    visio_headers = [
        "Process Step ID",
        "Process Step Description",
        "Next Step ID",
        "Connector Label",
        "Shape Type",
        "Function / Area",
    ]
    ciclo = [
        ("E00", "Lojista quer vender no JudgeTCG", "E01", "", "Start/End", "UI"),
        ("E01", "GET /entrar · Supabase Auth", "E02", "HTTPS", "External Data", "Supabase"),
        ("E02", "Sessão JWT?", "E03;E04", "Não;Sim", "Decision", "Auth"),
        ("E03", "Redirect /entrar?next=painel", "", "", "Start/End", "Auth"),
        ("E04", "GET /api/account/status KYC", "E05", "HTTP", "External Data", "FastAPI"),
        ("E05", "Status leu?", "E06;E07", "503;Sim", "Decision", "BFF"),
        ("E06", "Banner status indisponível (fail-open)", "E09", "", "Process", "UI"),
        ("E07", "can_publish / KYC ok?", "E08;E09", "Não;Sim", "Decision", "KYC"),
        ("E08", "GET /loja/suspensa", "", "", "Start/End", "KYC"),
        ("E09", "GET /api/stores/mine", "E10", "HTTP", "Data", "Stores"),
        ("E10", "Há loja?", "E11;E13", "Não;Sim", "Decision", "Stores"),
        ("E11", "POST /api/stores /stores/create + CNPJ", "E12;E13", "", "Subprocess", "Stores"),
        ("E12", "POST Stripe Connect onboard", "E13", "HTTPS", "External Data", "Stripe"),
        ("E13", "GET /api/seller/dashboard/overview", "E14", "HTTP", "Process", "Dashboard"),
        ("E14", "Overview 200?", "E15;E16", "Não;Sim", "Decision", "Dashboard"),
        ("E15", "Mock KPIs pending_payment=12 tickets=3", "E16", "", "Process", "Mock"),
        ("E16", "Catálogo GET cartas/produtos/jogos", "E17", "", "Subprocess", "Catálogo"),
        ("E17", "POST /api/seller/listings wizard", "E18;E19", "", "Subprocess", "Listagens"),
        ("E18", "POST /api/seller/inventory/*", "E19", "", "Subprocess", "Estoque"),
        ("E19", "Buyer vê PDP / listing", "E20", "", "Process", "Marketplace"),
        ("E20", "Checkout Stripe ou PIX", "E21", "HTTPS", "External Data", "Pagamento"),
        ("E21", "Pedido SQL marketplace", "E22", "SQL", "Data", "Orders"),
        ("E22", "GET /api/seller/orders + tabs", "E23;E24", "", "Subprocess", "Pedidos"),
        ("E23", "POST bulk/fulfillment/commands", "E25", "", "Data", "Fulfillment"),
        ("E24", "GET /api/seller/tickets", "E25", "", "Process", "Atendimento"),
        ("E25", "Repasse Stripe Connect / PIX", "E26", "", "Subprocess", "Financeiro"),
        ("E26", "LPC sem equipe JudgeTCG?", "E27;E28", "ainda não;quando houver evidência", "Decision", "North Star"),
        ("E27", "LPC = 0 Beta not started", "", "", "Start/End", "Produto"),
        ("E28", "Ciclo de produto desenhado", "", "", "Start/End", "Produto"),
    ]
    sh = wb.create_sheet("F01_Ciclo_E2E")
    sh.append(visio_headers)
    for r in ciclo:
        sh.append(list(r))
    for cell in sh[1]:
        cell.fill = fill
        cell.font = font
    for i, w in enumerate([16, 52, 18, 28, 16, 16], 1):
        sh.column_dimensions[get_column_letter(i)].width = w
    end = sh.max_row
    t = Table(displayName="F01CicloE2E", ref=f"A1:F{end}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sh.add_table(t)
    sh.freeze_panes = "A2"

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def main() -> None:
    pages = build_pages()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    vsdx = OUT_DIR / "JudgeTCG_Painel_Vendedor_Investidor.vsdx"
    write_vsdx(pages, vsdx)
    (OUT_DIR / "JudgeTCG_Painel_Vendedor_Investidor.drawio").write_text(drawio_xml(pages), encoding="utf-8")
    html = html_preview(pages).replace(
        "JudgeTCG — Fluxograma funcional E2E",
        "JudgeTCG — Painel do vendedor (investidor)",
    ).replace(
        "Arquivo MS Visio: docs/architecture/visio/JudgeTCG_Fluxograma_E2E.vsdx",
        "MS Visio: docs/architecture/visio/JudgeTCG_Painel_Vendedor_Investidor.vsdx",
    )
    (OUT_DIR / "JudgeTCG_Painel_Vendedor_Investidor.html").write_text(html, encoding="utf-8")
    xlsx = OUT_DIR / "JudgeTCG_Painel_Vendedor_Investidor.xlsx"
    write_xlsx(xlsx)
    print(f"Wrote {vsdx} ({vsdx.stat().st_size} bytes)")
    print(f"Wrote {xlsx} ({xlsx.stat().st_size} bytes)")
    print(f"Pages: {[p.name for p in pages]}")


if __name__ == "__main__":
    main()
